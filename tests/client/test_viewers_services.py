from __future__ import annotations

import hashlib
from pathlib import Path
import subprocess
from unittest.mock import patch

import openpyxl
from docx import Document
import pytest

from papagui_client.viewers import (
    ConversionOutcome,
    DocumentPreviewConverter,
    DocumentToolResolver,
    DocxTextReader,
    FilePreviewRouter,
    PreviewKind,
    SpreadsheetPreviewReader,
    TextFileReader,
)
from papagui_client.viewers.processes import PollingCommandRunner


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_file_router_and_bounded_text_reader_are_read_only(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("0123456789", encoding="utf-8")
    before = _digest(source)

    router = FilePreviewRouter()
    assert router.select(source).kind is PreviewKind.TEXT
    assert router.select(tmp_path / "book.pdf").kind is PreviewKind.PDF
    assert router.select(tmp_path / "table.xlsx").kind is PreviewKind.SPREADSHEET
    assert router.select(tmp_path / "letter.docx").kind is PreviewKind.WORD
    assert router.select(tmp_path / "photo.webp").kind is PreviewKind.IMAGE
    assert router.select(tmp_path / "archive.zip").kind is PreviewKind.UNSUPPORTED
    assert TextFileReader(maximum_characters=4).read(source) == "0123"
    assert _digest(source) == before


def test_docx_reader_includes_paragraphs_and_tables_without_mutation(tmp_path):
    source = tmp_path / "letter.docx"
    document = Document()
    document.add_paragraph("Hallo Kunde")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(source)
    before = _digest(source)

    assert DocxTextReader().read(source) == "Hallo Kunde\nA\tB"
    assert _digest(source) == before


def test_spreadsheet_reader_lists_sheets_and_bounds_preview(tmp_path):
    source = tmp_path / "table.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Übersicht"
    workbook.active.append(["Name", "Wert", "Extra"])
    workbook.active.append(["Muster", 42, "x"])
    workbook.create_sheet("Details")["A1"] = "zweites Blatt"
    workbook.save(source)
    before = _digest(source)

    reader = SpreadsheetPreviewReader(maximum_rows=1, maximum_columns=2)
    assert reader.sheet_names(source) == ("Übersicht", "Details")
    preview = reader.read_sheet(source, "Übersicht")

    assert preview.values == (("Name", "Wert"),)
    assert (preview.total_rows, preview.total_columns) == (2, 3)
    assert preview.truncated
    assert _digest(source) == before


def test_document_tool_resolver_maps_release_platforms_and_prefers_bundle(tmp_path):
    cases = (
        ("win32", "AMD64", "windows-x64"),
        ("darwin", "arm64", "macos-arm64"),
        ("darwin", "x86_64", "macos-x64"),
        ("linux", "x86_64", "linux-x64"),
    )
    for operating_system, machine, expected in cases:
        assert DocumentToolResolver.platform_tag(operating_system, machine) == expected

    binary = tmp_path / DocumentToolResolver.platform_tag() / "bin" / "pdftotext"
    binary.parent.mkdir(parents=True)
    binary.touch()
    resolver = DocumentToolResolver(tmp_path, path_lookup=lambda _name: "/fallback/tool")
    assert resolver.resolve("pdftotext") == binary

    binary.unlink()
    assert resolver.resolve("pdftotext") == Path("/fallback/tool")
    with pytest.raises(ValueError):
        resolver.resolve("../pdftotext")


class _Tools:
    def __init__(self, values: dict[str, Path] | None = None):
        self.values = values or {}

    def resolve(self, name: str) -> Path | None:
        return self.values.get(name)


class _Commands:
    def __init__(self, result: subprocess.CompletedProcess[str]):
        self.result = result
        self.calls = []

    def run(self, command, **options):
        self.calls.append((tuple(command), options))
        return self.result


def test_legacy_doc_uses_external_tool_then_binary_fallback(tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(
        "Unicode Zeile".encode("utf-16-le")
        + b"\x00\x01Readable ASCII line\r\n!!!!!\r\nReadable ASCII line"
    )
    successful = _Commands(subprocess.CompletedProcess([], 0, "external text", ""))
    converter = DocumentPreviewConverter(
        tmp_path / "cache",
        tools=_Tools({"catdoc": Path("/tools/catdoc")}),
        commands=successful,
    )
    external = converter.extract_legacy_doc(source)
    assert external.text == "external text"
    assert external.tool == "catdoc"

    converter = DocumentPreviewConverter(tmp_path / "cache", tools=_Tools())
    fallback = converter.extract_legacy_doc(source)
    assert "Readable ASCII line" in fallback.text
    assert "!!!!!" not in fallback.text
    assert fallback.tool == "Python (binär)"


def test_word_preview_cache_is_atomic_and_does_not_modify_source(tmp_path):
    source = tmp_path / "letter.docx"
    source.write_bytes(b"immutable source")
    rendered = tmp_path / "rendered.pdf"
    rendered.write_bytes(b"rendered pdf")
    converter = DocumentPreviewConverter(tmp_path / "cache", tools=_Tools())
    before = _digest(source)

    with (
        patch.object(converter, "_convert_with_libreoffice", return_value=rendered) as libre,
        patch.object(converter, "_convert_with_ms_office", return_value=None),
    ):
        first = converter.convert_word_to_pdf(source)
        second = converter.convert_word_to_pdf(source)

    assert first.path == second.path
    assert first.path is not None and first.path.read_bytes() == b"rendered pdf"
    assert first.tool == "LibreOffice"
    assert second.tool == "Cache"
    assert libre.call_count == 1
    assert _digest(source) == before
    assert not tuple((tmp_path / "cache").glob(".*.tmp"))


def test_conversion_outcomes_are_unambiguous_and_cancel_is_honoured(tmp_path):
    with pytest.raises(ValueError):
        ConversionOutcome("bad", "none", False)
    with pytest.raises(ValueError):
        ConversionOutcome("bad", "none", False, path=Path("x"), text="x")

    source = tmp_path / "sheet.xls"
    source.touch()
    converter = DocumentPreviewConverter(tmp_path / "cache", tools=_Tools())
    with (
        patch.object(converter, "_convert_with_libreoffice", return_value=None),
        pytest.raises(InterruptedError),
    ):
        converter.convert(source, "xlsx", should_cancel=lambda: True)


def test_polling_command_runner_returns_output_and_supports_cancel():
    runner = PollingCommandRunner()
    result = runner.run(
        ["/bin/sh", "-c", "printf viewer"],
        timeout=2,
    )
    assert result.returncode == 0
    assert result.stdout == "viewer"

    with pytest.raises(InterruptedError):
        runner.run(
            ["/bin/sh", "-c", "sleep 5"],
            timeout=2,
            should_cancel=lambda: True,
        )
