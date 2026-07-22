from pathlib import Path
from tempfile import TemporaryDirectory
import time
import unittest

import openpyxl
from docx import Document
from PyPDF2 import PdfWriter
from PySide6.QtCore import QtMsgType, qInstallMessageHandler
from PySide6.QtWidgets import QApplication
from unittest.mock import patch

from app.gui.viewer import FileViewer


class ViewerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def process_until(self, predicate, timeout: float = 5):
        deadline = time.monotonic() + timeout
        while not predicate() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(0.01)
        self.app.processEvents()

    def test_pdf_text_and_multisheet_excel_controls(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "test.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=400)
            with pdf_path.open("wb") as handle:
                writer.write(handle)

            excel_path = root / "test.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "Übersicht"
            workbook.active["A1"] = "Musterwert"
            workbook.create_sheet("Details")["B2"] = "Treffer"
            workbook.save(excel_path)

            text_path = root / "test.txt"
            text_path.write_text("eins zwei drei", encoding="utf-8")

            viewer = FileViewer()
            viewer.open_file(pdf_path)
            self.assertEqual(viewer.pdf_viewer.document.pageCount(), 1)
            viewer.open_file(excel_path)
            self.assertEqual(viewer.spreadsheet_viewer.sheet_combo.count(), 2)
            viewer.spreadsheet_viewer.sheet_combo.setCurrentText("Details")
            self.assertEqual(viewer.spreadsheet_viewer.table.item(1, 1).text(), "Treffer")
            viewer.open_file(text_path)
            viewer.text_viewer.search_input.setText("zwei")
            viewer.text_viewer.find_next()
            self.assertTrue(viewer.text_viewer.editor.textCursor().hasSelection())
            viewer.close()

    def test_spreadsheet_viewer_displays_explicit_truncation_hint(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            excel_path = root / "big.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Daten"
            for row in range(1, 505):
                for column in range(1, 52):
                    sheet.cell(row=row, column=column, value=f"{row}-{column}")
            workbook.save(excel_path)

            viewer = FileViewer()
            viewer.open_file(excel_path)

            self.assertIn("angezeigt bis 500 × 50", viewer.spreadsheet_viewer.info_label.text())
            self.assertIn("Hinweis: Anzeige ist auf 500 Zeilen und 50 Spalten begrenzt", viewer.spreadsheet_viewer.info_label.text())
            viewer.close()

    def test_pdf_close_document_does_not_emit_nullptr_warning(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "warncheck.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=400)
            with pdf_path.open("wb") as handle:
                writer.write(handle)

            messages: list[str] = []

            def collector(message_type, _context, message):
                if message_type in {QtMsgType.QtWarningMsg, QtMsgType.QtCriticalMsg}:
                    messages.append(str(message))

            previous_handler = qInstallMessageHandler(collector)
            try:
                viewer = FileViewer()
                viewer.open_file(pdf_path)
                baseline = sum(
                    "invalid nullptr parameter" in message for message in messages
                )
                viewer.pdf_viewer.close_document()
                self.app.processEvents()
                viewer.close()
            finally:
                qInstallMessageHandler(previous_handler)

            after_close = sum(
                "invalid nullptr parameter" in message for message in messages
            )
            self.assertEqual(after_close, baseline)

    def test_docx_preview_converts_to_pdf_asynchronously(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "preview.docx"
            document = Document()
            document.add_paragraph("Formatierter Inhalt")
            document.save(docx_path)
            pdf_path = root / "preview.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=400)
            with pdf_path.open("wb") as handle:
                writer.write(handle)

            def convert_to_pdf(*_args, **_kwargs):
                time.sleep(0.05)
                return pdf_path

            with patch(
                "app.services.document_converter.DocumentConverter.convert_word_to_pdf",
                side_effect=convert_to_pdf,
            ):
                viewer = FileViewer()
                viewer.open_file(docx_path)

                self.assertTrue(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.loading_widget)

                self.process_until(
                    lambda: viewer.stack.currentWidget() is viewer.pdf_viewer
                )
                self.process_until(lambda: not viewer.is_converting, timeout=1)

                self.assertFalse(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.pdf_viewer)
                self.assertEqual(viewer.pdf_viewer.document.pageCount(), 1)
                viewer.close()

    def test_docx_preview_falls_back_to_text_when_pdf_conversion_fails(self):
        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "fallback.docx"
            document = Document()
            document.add_paragraph("Fallback Inhalt")
            document.save(docx_path)

            with patch(
                "app.services.document_converter.DocumentConverter.convert_word_to_pdf",
                side_effect=RuntimeError("kein Konverter"),
            ):
                viewer = FileViewer()
                viewer.open_file(docx_path)

                self.process_until(
                    lambda: viewer.stack.currentWidget() is viewer.text_viewer
                )
                self.process_until(lambda: not viewer.is_converting, timeout=1)

                self.assertFalse(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.text_viewer)
                text = viewer.text_viewer.editor.toPlainText()
                self.assertIn("Formatierte Word-Vorschau konnte nicht erstellt werden", text)
                self.assertIn("Fallback Inhalt", text)
                viewer.close()

    def test_legacy_doc_processing_runs_asynchronously_with_loading_view(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.doc"
            path.write_bytes(
                b"Legacy customer document with enough readable text "
                b"for the Python fallback extraction."
            )

            with patch(
                "app.services.document_converter.DocumentConverter.convert_word_to_pdf",
                side_effect=RuntimeError("kein Konverter"),
            ):
                viewer = FileViewer()
                viewer.open_file(path)

                self.assertTrue(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.loading_widget)
                self.assertTrue(viewer.loading_indicator.is_running())

                self.process_until(
                    lambda: viewer.stack.currentWidget() is viewer.text_viewer
                )
                self.process_until(lambda: not viewer.is_converting, timeout=1)

                self.assertFalse(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.text_viewer)
                self.assertIn(
                    "Legacy customer document",
                    viewer.text_viewer.editor.toPlainText(),
                )
                viewer.close()


if __name__ == "__main__":
    unittest.main()
