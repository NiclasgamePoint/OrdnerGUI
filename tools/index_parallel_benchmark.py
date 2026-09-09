#!/usr/bin/env python3
"""Compare catalog workers using generated inputs in a disposable directory.

Run with the repository virtualenv, for example:
    .venv/bin/python tools/index_parallel_benchmark.py --mode controlled
    .venv/bin/python tools/index_parallel_benchmark.py --mode docx --documents 24

The controlled mode measures scheduling with fixed extraction latency. The DOCX
mode invokes the real parser on generated documents. Neither is a production
workload or a guarantee about customer indexing performance.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import threading
import time
from types import SimpleNamespace
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
for package in ("contracts", "server"):
    sys.path.insert(0, str(ROOT / "packages" / package / "src"))

from papagui_server.adapters.catalog import DocumentTextExtractor, SqliteCatalogIndexer  # noqa: E402
from papagui_server.domain.document_extraction import ExtractionResult  # noqa: E402
from papagui_server.domain.models import ServerSettings  # noqa: E402


class MeasuredExtractor(DocumentTextExtractor):
    def __init__(self, mode: str, latency: float) -> None:
        super().__init__()
        self.mode = mode
        self.latency = latency
        self.lock = threading.Lock()
        self.active = 0
        self.maximum_active = 0
        self.calls = 0

    def fingerprint(self, settings, extension=None):
        if self.mode == "controlled":
            return "synthetic-fixed-latency-v1"
        return super().fingerprint(settings, extension)

    def extract_document(self, path, settings, cancelled):
        with self.lock:
            self.active += 1
            self.maximum_active = max(self.maximum_active, self.active)
            self.calls += 1
        try:
            if self.mode == "controlled":
                time.sleep(self.latency)
                return ExtractionResult(text=path.read_text(encoding="utf-8"), status="ok")
            return super().extract_document(path, settings, cancelled)
        finally:
            with self.lock:
                self.active -= 1


def generate_documents(source: Path, mode: str, count: int) -> None:
    source.mkdir()
    for number in range(count):
        text = f"Generated document {number}. " + "Synthetic indexing benchmark. " * 30
        if mode == "docx":
            from docx import Document

            document = Document()
            document.add_paragraph(text)
            document.add_table(rows=1, cols=2).cell(0, 0).text = str(number)
            document.save(source / f"synthetic-{number:04}.docx")
        else:
            (source / f"synthetic-{number:04}.txt").write_text(text, encoding="utf-8")


def measure(source: Path, state: Path, workers: int, mode: str, latency: float) -> dict:
    extractor = MeasuredExtractor(mode, latency)
    indexer = SqliteCatalogIndexer(state, extractor)
    settings = ServerSettings(resource_profile="fast", ocr_enabled=False)
    original_open = Path.open
    opened = []

    def source_open(path, mode="r", *args, **kwargs):
        if source in path.parents and "r" in mode:
            opened.append(path.name)
        return original_open(path, mode, *args, **kwargs)

    runs = {}
    with (
        patch(
            "papagui_server.adapters.catalog.document_worker_budget",
            return_value=SimpleNamespace(workers=workers),
        ),
        patch.object(Path, "open", source_open),
    ):
        for kind in ("cold", "warm"):
            opened.clear()
            extractor.maximum_active = 0
            extractor.calls = 0
            start = time.perf_counter()
            indexer.build(
                source,
                source_id="synthetic-benchmark",
                full_rebuild=False,
                settings=settings,
                cancelled=lambda: False,
                progress=lambda *_: None,
            )
            elapsed = time.perf_counter() - start
            status = indexer.status()
            runs[kind] = {
                "elapsed_seconds": round(elapsed, 4),
                "maximum_concurrent_extractions": extractor.maximum_active,
                "extraction_calls": extractor.calls,
                "source_open_calls_in_coordinator_process": len(opened),
                "read_documents": status["read_documents"],
                "reused_documents": status["reused_documents"],
                "failed_documents": status["failed_documents"],
            }
    return {"workers": workers, **runs}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=("controlled", "docx"), default="controlled")
    parser.add_argument("--documents", type=int, default=24)
    parser.add_argument("--latency-ms", type=float, default=40)
    args = parser.parse_args()
    if args.documents < 1 or args.latency_ms < 0:
        parser.error("documents must be positive and latency must be nonnegative")
    with TemporaryDirectory(prefix="papagui-synthetic-index-") as temporary:
        root = Path(temporary)
        source = root / "generated-source"
        generate_documents(source, args.mode, args.documents)
        results = [
            measure(
                source,
                root / f"generated-state-{workers}",
                workers,
                args.mode,
                args.latency_ms / 1000,
            )
            for workers in (1, 4)
        ]
    serial_seconds = results[0]["cold"]["elapsed_seconds"]
    parallel_seconds = results[1]["cold"]["elapsed_seconds"]
    print(
        json.dumps(
            {
                "scope": "Synthetic inputs only; not a production performance claim.",
                "mode": args.mode,
                "documents": args.documents,
                "controlled_latency_ms": args.latency_ms if args.mode == "controlled" else None,
                "cold_elapsed_ratio_serial_to_parallel": round(serial_seconds / parallel_seconds, 3)
                if parallel_seconds
                else None,
                "runs": results,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
