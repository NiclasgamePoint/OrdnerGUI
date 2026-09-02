"""Read-only spreadsheet preview adapter with bounded memory use."""

from __future__ import annotations

from pathlib import Path

from .models import SheetPreview


class SpreadsheetPreviewReader:
    def __init__(self, maximum_rows: int = 500, maximum_columns: int = 50) -> None:
        if maximum_rows < 1 or maximum_columns < 1:
            raise ValueError("spreadsheet preview limits must be positive")
        self.maximum_rows = maximum_rows
        self.maximum_columns = maximum_columns

    def sheet_names(self, path: Path) -> tuple[str, ...]:
        if path.suffix.lower() == ".xlsx":
            workbook = self._open_xlsx(path)
            try:
                return tuple(workbook.sheetnames)
            finally:
                workbook.close()
        workbook = self._open_xls(path)
        try:
            return tuple(workbook.sheet_names())
        finally:
            workbook.release_resources()

    def read_sheet(self, path: Path, sheet_name: str) -> SheetPreview:
        if path.suffix.lower() == ".xlsx":
            return self._read_xlsx(path, sheet_name)
        return self._read_xls(path, sheet_name)

    @staticmethod
    def _open_xlsx(path: Path):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - depends on optional install
            raise RuntimeError("openpyxl ist für XLSX-Vorschauen erforderlich") from exc
        return openpyxl.load_workbook(str(path), read_only=True, data_only=True)

    @staticmethod
    def _open_xls(path: Path):
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - depends on optional install
            raise RuntimeError("xlrd ist für XLS-Vorschauen erforderlich") from exc
        return xlrd.open_workbook(str(path), on_demand=True)

    def _read_xlsx(self, path: Path, sheet_name: str) -> SheetPreview:
        workbook = self._open_xlsx(path)
        try:
            sheet = workbook[sheet_name]
            total_rows = sheet.max_row or 0
            total_columns = sheet.max_column or 0
            rows = min(self.maximum_rows, total_rows)
            columns = min(self.maximum_columns, total_columns)
            values = tuple(
                tuple(row)
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=rows,
                    min_col=1,
                    max_col=columns,
                    values_only=True,
                )
            ) if rows and columns else ()
        finally:
            workbook.close()
        return SheetPreview(
            name=sheet_name,
            values=values,
            total_rows=total_rows,
            total_columns=total_columns,
            maximum_rows=self.maximum_rows,
            maximum_columns=self.maximum_columns,
        )

    def _read_xls(self, path: Path, sheet_name: str) -> SheetPreview:
        workbook = self._open_xls(path)
        try:
            sheet = workbook.sheet_by_name(sheet_name)
            total_rows, total_columns = sheet.nrows, sheet.ncols
            rows = min(self.maximum_rows, total_rows)
            columns = min(self.maximum_columns, total_columns)
            values = tuple(
                tuple(sheet.row_values(row, 0, columns)) for row in range(rows)
            )
        finally:
            workbook.release_resources()
        return SheetPreview(
            name=sheet_name,
            values=values,
            total_rows=total_rows,
            total_columns=total_columns,
            maximum_rows=self.maximum_rows,
            maximum_columns=self.maximum_columns,
        )
