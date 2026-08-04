from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
import os
import shutil
import subprocess
import sys
import time
from tempfile import TemporaryDirectory

import openpyxl
from docx import Document
from PyPDF2 import PdfReader

from app.core.config import IndexOptions


class DocumentTextIndexer:
    """Extract bounded searchable text without knowing any index persistence."""

    LEGACY_XLS_TIMEOUT_SECONDS = 15
    LEGACY_DOC_TIMEOUT_SECONDS = 20
    TEXT_TYPES = {"txt", "csv", "md", "log", "json", "xml", "yaml", "yml", "ini"}

    def __init__(self, options: IndexOptions):
        self.options = options
        self._ocr_language: str | None = None

    def extract(self, path: Path) -> tuple[str, str, str]:
        maximum = self.options.max_file_size_mb * 1024 * 1024
        try:
            if path.stat().st_size > maximum:
                return (
                    "",
                    "skipped_large",
                    f"Datei größer als {self.options.max_file_size_mb} MB",
                )
        except OSError as exc:
            return "", "error", str(exc)
        try:
            text = self._extract(path, path.suffix.lower().lstrip("."))
            return text, ("success" if text.strip() else "empty"), ""
        except TimeoutError as exc:
            return "", "timeout", str(exc)
        except Exception as exc:
            message = str(exc)
            status = (
                "encrypted"
                if "encrypted" in message.casefold() or "password" in message.casefold()
                else "error"
            )
            return "", status, message

    def _extract(self, path: Path, file_type: str) -> str:
        if file_type in self.TEXT_TYPES:
            return self._limit(path.read_text(encoding="utf-8", errors="replace"))
        if file_type == "pdf":
            return self._pdf(path)
        if file_type == "docx":
            return self._docx(path)
        if file_type == "doc":
            return self._legacy(path, "doc")
        if file_type == "xlsx":
            return self._xlsx(path)
        if file_type == "xls":
            return self._legacy(path, "xls")
        return ""

    def _limit(self, text: str) -> str:
        return text[: self.options.max_extracted_characters]

    def _pdf(self, path: Path) -> str:
        reader = PdfReader(str(path))
        parts: list[str] = []
        length = 0
        for page in reader.pages:
            text = page.extract_text() or ""
            if text:
                parts.append(text)
                length += len(text)
            if length >= self.options.max_extracted_characters:
                break
        extracted = self._limit("\n".join(parts))
        if extracted.strip() or not self.options.ocr_enabled:
            return extracted
        return self._ocr_pdf(path)

    def _docx(self, path: Path) -> str:
        document = Document(str(path))
        parts = [item.text for item in document.paragraphs if item.text]
        length = sum(map(len, parts))
        for table in document.tables:
            for row in table.rows:
                values = [cell.text for cell in row.cells if cell.text]
                if values:
                    line = "\t".join(values)
                    parts.append(line)
                    length += len(line)
                if length >= self.options.max_extracted_characters:
                    return self._limit("\n".join(parts))
        return self._limit("\n".join(parts))

    def _xlsx(self, path: Path) -> str:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            parts: list[str] = []
            length = 0
            for worksheet in workbook.worksheets:
                parts.append(worksheet.title)
                for row in worksheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        line = "\t".join(values)
                        parts.append(line)
                        length += len(line)
                    if length >= self.options.max_extracted_characters:
                        return self._limit("\n".join(parts))
            return self._limit("\n".join(parts))
        finally:
            workbook.close()

    def _legacy(self, path: Path, file_type: str) -> str:
        module = f"app.services.{file_type}_text_extractor"
        timeout = (
            self.LEGACY_XLS_TIMEOUT_SECONDS
            if file_type == "xls" else self.LEGACY_DOC_TIMEOUT_SECONDS
        )
        options = {
            "capture_output": True,
            "text": True,
            "encoding": "utf-8",
            "errors": "replace",
            "timeout": timeout,
            "cwd": str(Path(__file__).resolve().parents[2]),
        }
        if sys.platform == "win32":
            options["creationflags"] = subprocess.CREATE_NO_WINDOW
        try:
            result = subprocess.run(
                [
                    sys.executable, "-m", module, str(path),
                    "--maximum-characters", str(self.options.max_extracted_characters),
                ],
                **options,
            )
        except subprocess.TimeoutExpired as exc:
            raise TimeoutError(
                f"{file_type.upper()}-Zeitlimit von {timeout} Sekunden erreicht"
            ) from exc
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or f"Unbekannter {file_type.upper()}-Lesefehler")
        return self._limit(result.stdout)

    def _ocr_pdf(self, path: Path) -> str:
        pdftoppm = shutil.which("pdftoppm")
        tesseract = self._find_tesseract()
        if pdftoppm is None or tesseract is None:
            return ""
        deadline = time.monotonic() + self.options.ocr_timeout_seconds
        with TemporaryDirectory(prefix="papagui-ocr-") as directory:
            prefix = Path(directory) / "page"
            render = subprocess.run(
                [
                    pdftoppm, "-jpeg", "-jpegopt", "quality=85", "-r", "120",
                    "-f", "1", "-l", str(self.options.ocr_max_pages),
                    str(path), str(prefix),
                ],
                capture_output=True,
                text=True,
                timeout=max(5, min(15, self.options.ocr_timeout_seconds)),
            )
            if render.returncode != 0:
                return ""
            images = sorted(Path(directory).glob("page-*.jpg"))
            language = self._get_ocr_language(tesseract)

            def recognize(image: Path) -> tuple[Path, str]:
                remaining = max(1, deadline - time.monotonic())
                result = subprocess.run(
                    [str(tesseract), str(image), "stdout", "-l", language],
                    capture_output=True,
                    text=True,
                    errors="replace",
                    timeout=max(1, min(15, remaining)),
                )
                return image, result.stdout if result.returncode == 0 else ""

            recognized: dict[Path, str] = {}
            with ThreadPoolExecutor(max_workers=min(4, max(1, len(images)))) as executor:
                futures = [executor.submit(recognize, image) for image in images]
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise TimeoutError(
                        f"OCR-Zeitlimit von {self.options.ocr_timeout_seconds} Sekunden erreicht"
                    )
                try:
                    for future in as_completed(futures, timeout=remaining):
                        image, text = future.result()
                        if text.strip():
                            recognized[image] = text
                except TimeoutError as exc:
                    raise TimeoutError(
                        f"OCR-Zeitlimit von {self.options.ocr_timeout_seconds} Sekunden erreicht"
                    ) from exc
            return self._limit("\n".join(recognized[image] for image in images if image in recognized))

    def _get_ocr_language(self, tesseract: Path) -> str:
        if self._ocr_language is not None:
            return self._ocr_language
        try:
            result = subprocess.run(
                [str(tesseract), "--list-langs"], capture_output=True, text=True, timeout=10
            )
            languages = set(result.stdout.splitlines()[1:])
        except Exception:
            languages = set()
        self._ocr_language = (
            "deu+eng" if {"deu", "eng"}.issubset(languages)
            else ("deu" if "deu" in languages else "eng")
        )
        return self._ocr_language

    @staticmethod
    def _find_tesseract() -> Path | None:
        found = shutil.which("tesseract")
        if found:
            return Path(found)
        if sys.platform == "win32":
            for name in ("PROGRAMFILES", "LOCALAPPDATA"):
                base = os.environ.get(name)
                if base:
                    candidate = Path(base) / "Tesseract-OCR" / "tesseract.exe"
                    if candidate.exists():
                        return candidate
        return None
