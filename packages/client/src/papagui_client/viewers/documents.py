"""Read-only routing and extraction for directly displayable documents."""

from __future__ import annotations

from pathlib import Path

from .models import PreviewKind, PreviewSelection


class FilePreviewRouter:
    TEXT_TYPES = frozenset({"txt", "csv", "log", "md", "json", "xml", "yaml", "yml", "ini"})
    IMAGE_TYPES = frozenset({"jpg", "jpeg", "png", "gif", "bmp", "webp"})

    def select(self, path: Path) -> PreviewSelection:
        source = path.resolve()
        suffix = source.suffix.lower().lstrip(".")
        if suffix == "pdf":
            kind = PreviewKind.PDF
        elif suffix in {"xls", "xlsx"}:
            kind = PreviewKind.SPREADSHEET
        elif suffix in self.TEXT_TYPES:
            kind = PreviewKind.TEXT
        elif suffix in {"doc", "docx"}:
            kind = PreviewKind.WORD
        elif suffix in self.IMAGE_TYPES:
            kind = PreviewKind.IMAGE
        else:
            kind = PreviewKind.UNSUPPORTED
        return PreviewSelection(path=source, kind=kind)


class TextFileReader:
    def __init__(self, maximum_characters: int = 2_000_000) -> None:
        if maximum_characters < 1:
            raise ValueError("maximum_characters must be positive")
        self.maximum_characters = maximum_characters

    def read(self, path: Path) -> str:
        with path.open("r", encoding="utf-8", errors="replace") as source:
            return source.read(self.maximum_characters)


class DocxTextReader:
    """Extract paragraphs and table cells without modifying the document."""

    def read(self, path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - depends on optional install
            raise RuntimeError("python-docx ist für die DOCX-Textvorschau erforderlich") from exc

        document = Document(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(lines)
