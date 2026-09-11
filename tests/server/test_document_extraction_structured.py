"""Only generated, synthetic documents are used by these parser regressions."""

from __future__ import annotations

from dataclasses import replace
import json
import os
from pathlib import Path
import time

import pytest

from papagui_server.adapters.catalog import SqliteCatalogIndexer
from papagui_server.adapters.catalog_extraction import (
    DocumentTextExtractor,
    ExternalCommandRunner,
    extraction_fingerprint,
)
from papagui_server.adapters.catalog_reader import SqliteCatalogReader
from papagui_server.adapters.extraction_office import docx, xlsx, display_value
from papagui_server.adapters.extraction_ocr import tsv_blocks
from papagui_server.adapters.extraction_pdf_worker import read_pdf
from papagui_server.adapters.extraction_store import ExtractionStore
from papagui_server.domain.document_extraction import (
    ExtractionBlock,
    ExtractionResult,
    bounded_result,
)
from papagui_server.domain.models import ServerSettings


def test_structured_result_roundtrip_and_single_budget() -> None:
    block = ExtractionBlock("abcdefgh", page=2, bbox=(1, 2, 3, 4), confidence=0.8)
    result = bounded_result([block, ExtractionBlock("ijk")], 6)
    assert (result.text, result.status, result.reason) == ("abcdef", "partial", "character_budget")
    assert result.blocks[0].bbox == (1, 2, 3, 4)
    assert ExtractionResult.from_dict(json.loads(json.dumps(result.to_dict()))) == result
    assert bounded_result([ExtractionBlock(" ")], 20).status == "no_text"
    assert bounded_result([block, block], 8).status == "partial"
    with pytest.raises(ValueError):
        ExtractionResult(status="not_a_status")


def _word_fixture(path: Path) -> None:
    from docx import Document

    document = Document()
    header = document.sections[0].header.paragraphs[0]
    header.add_run("Absender: ")
    header.add_run("Synthetisches Büro")
    paragraph = document.add_paragraph()
    paragraph.add_run("Auftrag")
    paragraph.add_run("geber: Beispiel GmbH")
    paragraph.add_run("\nFantasiestraße 12\n01234 Teststadt")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Telefon"
    table.cell(0, 1).text = "+49 30 1234567"
    document.sections[0].footer.paragraphs[0].text = "Synthetischer Fußtext"
    document.save(path)


def test_docx_joins_runs_and_keeps_headers_footers_and_table_rows(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.docx"
    _word_fixture(path)
    result = docx(path, 20_000, lambda: False, time.monotonic() + 10)
    assert result.status == "ok"
    assert "Auftraggeber: Beispiel GmbH\nFantasiestraße" in result.text
    assert "Auftrag\ngeber" not in result.text
    assert result.blocks[0].section == "header1"
    assert result.blocks[-1].section == "footer1"
    assert any(
        block.kind == "table_row" and block.text == "Telefon\t+49 30 1234567"
        for block in result.blocks
    )
    assert docx(path, 4, lambda: False, time.monotonic() + 10).status == "partial"
    assert docx(path, 20_000, lambda: True, time.monotonic() + 10).reason == "cancelled"
    assert docx(path, 20_000, lambda: False, 0).reason == "time_budget"
    assert DocumentTextExtractor().extract_document(path, ServerSettings()).text == result.text


def _spreadsheet(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Kontakte"
    sheet.append(["Beispiel GmbH", "PLZ", "Telefon"])
    sheet.append(["Synthetische Person", 1234, 301234567])
    sheet["B2"].number_format = "00000"
    sheet["C2"].number_format = "0000000000"
    sheet["A3"] = "=1+2"
    workbook.save(path)
    workbook.close()


def test_xlsx_correct_values_coordinates_leading_zero_and_missing_formula(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.xlsx"
    _spreadsheet(path)
    result = xlsx(path, 20_000, lambda: False, time.monotonic() + 10)
    values = {block.cell: block.text for block in result.blocks}
    assert values["B2"] == "01234"
    assert values["C2"] == "0301234567"
    assert "=1+2" not in result.text
    assert {block.sheet for block in result.blocks} == {"Kontakte"}
    assert (result.status, result.reason) == ("partial", "formula_result_missing")
    assert xlsx(path, 4, lambda: False, time.monotonic() + 10).reason == "character_budget"
    assert xlsx(path, 200, lambda: True, time.monotonic() + 10).reason == "cancelled"
    assert DocumentTextExtractor().extract_document(path, ServerSettings()).text == result.text


@pytest.mark.parametrize(
    "value,mask,expected",
    [
        (1234, "00000", "01234"),
        (True, "", "TRUE"),
        (None, "", ""),
        (1.25, "0.00", "1.25"),
        (12, "000 000", "000 012"),
        (1234, "0", "1234"),
        (-12, "0000", "-0012"),
        ("001", "", "001"),
    ],
)
def test_spreadsheet_display_values(value, mask, expected) -> None:
    assert display_value(value, mask) == expected


def test_tsv_keeps_paragraph_lines_boxes_and_confidence() -> None:
    data = b"level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n5\t1\t1\t1\t1\t1\t10\t20\t30\t10\t80\tBeispiel\n5\t1\t1\t1\t2\t1\t10\t40\t50\t10\t100\tTelefon\n5\t1\t2\t1\t1\t1\tbad\t0\t0\t0\t0\tignored\n"
    blocks = tsv_blocks(data, page=3, scale=0.5)
    assert len(blocks) == 1
    assert blocks[0].text == "Beispiel\nTelefon"
    assert blocks[0].bbox == (5, 10, 30, 25)
    assert blocks[0].confidence == 0.9
    assert blocks[0].page == 3
    assert tsv_blocks(b"", page=1) == []
    assert tsv_blocks(b"injected text", page=1)[0].text == "injected text"


def _pdf(path: Path, pages: int = 3) -> None:
    from reportlab.pdfgen.canvas import Canvas

    canvas = Canvas(str(path))
    canvas.drawString(40, 750, "Auftraggeber Beispiel GmbH " + "Synthetischer Kontakt " * 5)
    canvas.drawString(40, 735, "Fantasiestrasse 12")
    canvas.drawString(40, 720, "01234 Teststadt")
    canvas.showPage()
    for _ in range(pages - 1):
        canvas.showPage()
    canvas.save()


def test_real_pdf_layout_page_budget_and_isolated_parser(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.pdf"
    _pdf(path)
    native = read_pdf(path, 10_000, 2)
    assert (native.status, native.reason, native.pages_total, native.pages_processed) == (
        "partial",
        "page_budget",
        3,
        2,
    )
    assert native.blocks[0].page == 1 and native.blocks[0].bbox
    assert "Fantasiestrasse 12\n01234 Teststadt" in native.text
    assert native.pages[1].status == "no_text"
    assert read_pdf(path, 10, 3).reason == "character_budget"
    actual = DocumentTextExtractor().extract_document(path, ServerSettings(ocr_enabled=False))
    assert actual.status == "partial" and actual.reason == "ocr_disabled"
    assert actual.pages_total == 3


class PageRunner(ExternalCommandRunner):
    def __init__(self):
        super().__init__()
        self.calls = []

    def run(self, command, *, timeout):
        self.calls.append(command)
        self.last_status = "ok"
        if command[0] == "pdftoppm":
            from PIL import Image

            Image.new("RGB", (100, 200), "white").save(Path(command[-1]).with_suffix(".png"))
            return b""
        return b"OCR Auftraggeber Beispiel GmbH\nTelefon +49 30 1234567"


def test_mixed_pdf_ocr_is_page_local_and_budget_visible(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.pdf"
    _pdf(path)
    runner = PageRunner()
    extractor = DocumentTextExtractor(runner)
    result = extractor.extract_document(
        path, ServerSettings(ocr_max_pages=1, ocr_extension_threshold=100)
    )
    assert result.status == "partial"
    assert result.reason == "ocr_page_budget"
    assert result.text.count("Synthetischer Kontakt") == 5
    assert result.text.count("OCR Auftraggeber") == 1
    assert [page.status for page in result.pages] == ["ok", "ok", "partial"]
    assert len([command for command in runner.calls if command[0] == "tesseract"]) == 1
    assert runner.calls[0][runner.calls[0].index("-f") + 1] == "2"


def test_image_probe_ocr_and_pixel_limit(tmp_path: Path) -> None:
    from PIL import Image

    extractor = DocumentTextExtractor(PageRunner())
    scan = tmp_path / "scan.png"
    Image.new("RGB", (1200, 1200), "white").save(scan)
    result = extractor.extract_document(scan, ServerSettings())
    assert result.status == "ok" and result.pages_processed == 1
    assert result.blocks[0].page == 1
    assert (
        extractor.extract_document(scan, ServerSettings(ocr_enabled=False)).reason == "ocr_disabled"
    )
    small_budget = ServerSettings(image_max_pixels=1000000)
    assert extractor.extract_document(scan, small_budget).reason == "image_pixel_budget"
    photo = tmp_path / "photo.png"
    Image.new("RGB", (200, 200), "blue").save(photo)
    result = extractor.extract_document(photo, ServerSettings())
    assert (result.status, result.reason) == ("partial", "image_not_document")


def test_timeout_tool_failures_and_malformed_files_are_sanitized(
    tmp_path: Path, monkeypatch
) -> None:
    extractor = DocumentTextExtractor()
    secret_path = tmp_path / "synthetic-private-name.pdf"
    secret_path.write_bytes(b"invalid")
    monkeypatch.setattr(extractor._worker_runner, "run", lambda *_args, **_kwargs: b"")
    extractor._worker_runner.last_status = "timeout"
    result = extractor.extract_document(secret_path, ServerSettings())
    assert result.status == "timeout"
    assert "synthetic-private-name" not in json.dumps(result.to_dict())
    assert DocumentTextExtractor._remaining(time.monotonic() + 20, 3) == 3
    with pytest.raises(TimeoutError):
        DocumentTextExtractor._remaining(0, 3)
    assert (
        extractor.extract_document(tmp_path / "missing.bin", ServerSettings()).status
        == "unsupported"
    )
    assert extractor.extract_document(tmp_path / "missing.txt", ServerSettings()).status == "error"
    secret_path.write_bytes(b"x" * (1024 * 1024 + 1))
    assert (
        extractor.extract_document(secret_path, ServerSettings(max_file_size_mb=1)).status
        == "too_large"
    )


def test_extraction_cache_retries_versions_and_protected_cleanup(
    tmp_path: Path, monkeypatch
) -> None:
    store = ExtractionStore(tmp_path / "private" / "artifacts.db")
    assert store.get("hash", "version") is None
    assert store.read("hash", "version") is None
    result = ExtractionResult(
        text="synthetic",
        status="error",
        reason="temporary",
        content_hash="hash",
        version_hash="version",
    )
    assert store.put(result, retry_delay_seconds=100)
    assert store.get("hash", "version") == result
    assert store.read("hash", "version") == result
    assert store.get("hash", "different-version") is None
    now = time.time()
    monkeypatch.setattr("papagui_server.adapters.extraction_store.time.time", lambda: now + 200)
    assert store.get("hash", "version") is None
    assert store.get("hash", "version", max_attempts=1) == result
    assert store.put(result, retry_delay_seconds=0)
    assert store.put(result, retry_delay_seconds=0)
    assert store.get("hash", "version") == result
    monkeypatch.setattr(
        "papagui_server.adapters.extraction_store.time.time", lambda: now + 90 * 86400
    )
    assert store.cleanup(protected=[("hash", "version")], retention_days=30) == 0
    assert store.cleanup(protected=[], retention_days=30) == 1
    assert not store.put(result, max_store_mb=0)


class CountStructured:
    def __init__(self):
        self.calls = 0
        self.version = "version-1"

    def fingerprint(self, _settings):
        return self.version

    def extract_document(self, path, settings, cancelled):
        self.calls += 1
        text = path.read_text()
        return ExtractionResult(
            text=text, blocks=(ExtractionBlock(text, section="synthetic"),), status="ok"
        )


def _build(indexer, source, **kwargs):
    return indexer.build(
        source,
        source_id="primary",
        full_rebuild=kwargs.pop("full_rebuild", True),
        settings=kwargs.pop("settings", ServerSettings()),
        cancelled=lambda: False,
        progress=lambda *_: None,
        **kwargs,
    )


def test_rebuild_cache_copy_metadata_change_and_parser_invalidation(tmp_path: Path) -> None:
    project = tmp_path / "source" / "Planung" / "2026" / "Beispiel GmbH"
    project.mkdir(parents=True)
    document = project / "Anschreiben.txt"
    document.write_text("synthetic-alpha")
    extractor = CountStructured()
    indexer = SqliteCatalogIndexer(tmp_path / "server", extractor)
    _build(indexer, tmp_path / "source")
    _build(indexer, tmp_path / "source")
    assert extractor.calls == 1
    (project / "Kopie.txt").write_bytes(document.read_bytes())
    _build(indexer, tmp_path / "source")
    assert extractor.calls == 1  # Same first-page policy and content share the cache.
    before = document.stat()
    document.write_text("synthetic-bravo")
    os.utime(document, ns=(before.st_atime_ns, before.st_mtime_ns))
    _build(indexer, tmp_path / "source", full_rebuild=False)
    assert extractor.calls == 1  # Ordinary scans trust unchanged modification time and size.
    _build(indexer, tmp_path / "source", full_rebuild=False, verify_content=True)
    assert extractor.calls == 2
    extractor.version = "version-2"
    _build(indexer, tmp_path / "source", full_rebuild=False)
    assert extractor.calls == 4
    _build(indexer, tmp_path / "source", full_rebuild=False, force_extraction=True)
    assert extractor.calls == 6
    reader = SqliteCatalogReader(indexer.active_path)
    evidence = reader.document_evidence(source_id="primary")
    assert all(item["blocks"] and item["content_hash"] for item in evidence)
    assert indexer.extraction_store.database_path.parent.parent == tmp_path / "server"
    assert "index" not in indexer.extraction_store.database_path.parts
    coverage = reader.coverage(source_id="primary")[0]
    assert (coverage["files_total"], coverage["files_extracted"], coverage["status_counts"]) == (
        2,
        2,
        {"ok": 2},
    )


def test_reader_second_round_priority_source_paths_and_coverage(tmp_path: Path) -> None:
    source = tmp_path / "source"
    root = source / "Planung" / "2026" / "Beispiel GmbH"
    root.mkdir(parents=True)
    for number in range(30):
        (root / f"Berechnung-{number:02}.txt").write_text("Mengen Statik " + str(number))
    (root / "Altes Anschreiben.txt").write_text(
        "Auftraggeber Beispiel GmbH\nTelefon +49 30 1234567"
    )
    (root / "Unbekannt.bin").write_bytes(b"synthetic")
    indexer = SqliteCatalogIndexer(tmp_path / "server")
    _build(indexer, source)
    reader = SqliteCatalogReader(indexer.active_path)
    first = reader.document_evidence(source_id="primary", documents_per_project=24)
    assert first[0]["source"]["relative_path"].endswith("Altes Anschreiben.txt")
    second = reader.document_evidence(
        source_id="primary", documents_per_project=24, offset_per_project=24
    )
    assert len(first) == 24 and len(second) == 7
    assert not (
        {item["source"]["relative_path"] for item in first}
        & {item["source"]["relative_path"] for item in second}
    )
    assert len(reader.source_paths(source_id="primary")) == 32
    coverage = reader.coverage(source_id="primary")[0]
    assert coverage["status_counts"] == {"ok": 31, "unsupported": 1}
    assert coverage["reason_counts"] == {"unsupported_extension": 1}
    assert reader.document_evidence(source_id="primary", project_root_ids=[]) == []
    assert reader.coverage(source_id="primary", project_root_ids=[]) == []
    assert reader.source_paths(source_id="primary", project_root_ids=[]) == []
    with pytest.raises(ValueError):
        reader.coverage(source_id="primary", project_root_ids=[True])
    with pytest.raises(ValueError):
        reader.document_evidence(source_id="primary", offset_per_project=-1)


def test_relevant_settings_and_tool_versions_invalidate_fingerprint() -> None:
    settings = ServerSettings()
    assert extraction_fingerprint(settings) != extraction_fingerprint(
        replace(settings, ocr_max_pages=3)
    )
    assert extraction_fingerprint(settings) == extraction_fingerprint(
        replace(settings, interval_seconds=1800)
    )


def test_real_biff_xls_preserves_cells_dates_and_numeric_masks(tmp_path: Path) -> None:
    import xlwt
    from datetime import datetime
    from papagui_server.adapters.extraction_office import xls

    path = tmp_path / "synthetic.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Kontaktblatt")
    sheet.write(0, 0, "Beispiel GmbH")
    sheet.write(1, 1, 1234, xlwt.easyxf(num_format_str="00000"))
    sheet.write(1, 2, 301234567, xlwt.easyxf(num_format_str="0000000000"))
    sheet.write(2, 0, datetime(2026, 1, 2), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
    workbook.save(str(path))
    result = xls(path, 2000, lambda: False, time.monotonic() + 10)
    values = {block.cell: block.text for block in result.blocks}
    assert values["B2"] == "01234"
    assert values["C2"] == "0301234567"
    assert values["A3"].startswith("2026-01-02")
    assert {block.sheet for block in result.blocks} == {"Kontaktblatt"}
    assert DocumentTextExtractor().extract_document(path, ServerSettings()).text == result.text
    assert xls(path, 2000, lambda: True, time.monotonic() + 10).reason == "cancelled"
    assert xls(path, 5, lambda: False, time.monotonic() + 10).reason == "character_budget"


def test_targeted_force_retains_other_project_cache_and_pins_snapshot(tmp_path: Path) -> None:
    source = tmp_path / "source"
    for name in ("Beispiel A", "Beispiel B"):
        root = source / "Planung" / "2026" / name
        root.mkdir(parents=True)
        (root / "Kontakt.txt").write_text("synthetic " + name)
    extractor = CountStructured()
    indexer = SqliteCatalogIndexer(tmp_path / "server", extractor)
    _build(indexer, source)
    reader = SqliteCatalogReader(indexer.active_path)
    project_id = reader.list_project_roots(source_id="primary")[0]["id"]
    _build(
        indexer, source, full_rebuild=False, force_extraction=True, project_root_ids=[project_id]
    )
    assert extractor.calls == 3
    with reader.snapshot() as snapshot:
        pinned_version = snapshot.snapshot_version()
        (source / "Planung" / "2026" / "Beispiel A" / "Kontakt.txt").unlink()
        _build(indexer, source)
        assert len(snapshot.source_paths(source_id="primary")) == 2
        assert len(reader.source_paths(source_id="primary")) == 1
        assert snapshot.snapshot_version() == pinned_version
        assert all(item["blocks"] for item in snapshot.document_evidence(source_id="primary"))
        snapshot_path = snapshot.database_path
    assert not snapshot_path.exists()


def test_worker_error_categories_without_leaking_exception_messages(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    import sys
    from types import SimpleNamespace
    from papagui_server.adapters import extraction_pdf_worker, extraction_office_worker

    monkeypatch.setitem(sys.modules, "resource", SimpleNamespace(RLIMIT_AS=9, setrlimit=lambda *_: None))
    monkeypatch.setattr(sys, "argv", ["worker", str(tmp_path / "synthetic.docx"), "100", "2"])
    for exception, status in (
        (ModuleNotFoundError("synthetic-private"), "tool_missing"),
        (MemoryError("synthetic-private"), "too_large"),
        (ValueError("synthetic-private"), "error"),
    ):

        def fail(*_, error=exception, **__):
            raise error

        monkeypatch.setattr(extraction_office_worker.extraction_office, "docx", fail)
        extraction_office_worker.main()
        payload = capsys.readouterr().out
        assert json.loads(payload)["status"] == status
        assert "synthetic-private" not in payload
    monkeypatch.setattr(
        extraction_office_worker.extraction_office,
        "docx",
        lambda *_: ExtractionResult(text="synthetic", status="ok"),
    )
    extraction_office_worker.main()
    assert json.loads(capsys.readouterr().out)["status"] == "ok"

    class PDFPasswordIncorrect(Exception):
        pass

    for exception, status in (
        (ModuleNotFoundError(), "tool_missing"),
        (MemoryError(), "too_large"),
        (PDFPasswordIncorrect(), "encrypted"),
        (ValueError("private"), "error"),
    ):

        def fail(*_, error=exception):
            raise error

        monkeypatch.setattr(extraction_pdf_worker, "read_pdf", fail)
        extraction_pdf_worker.main()
        assert json.loads(capsys.readouterr().out)["status"] == status
    monkeypatch.setattr(
        extraction_pdf_worker, "read_pdf", lambda *_: ExtractionResult(status="no_text")
    )
    extraction_pdf_worker.main()
    assert json.loads(capsys.readouterr().out)["status"] == "no_text"


def test_office_archive_and_time_cell_limits(tmp_path: Path, monkeypatch) -> None:
    from papagui_server.adapters import extraction_office

    encrypted = tmp_path / "encrypted.docx"
    encrypted.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    with pytest.raises(ValueError, match="encrypted_office"):
        extraction_office.check_office_archive(encrypted)
    assert (
        DocumentTextExtractor().extract_document(encrypted, ServerSettings()).status == "encrypted"
    )
    path = tmp_path / "synthetic.xlsx"
    _spreadsheet(path)
    assert xlsx(path, 10000, lambda: False, 0).reason == "time_budget"
    monkeypatch.setattr(extraction_office, "MAX_CELLS", 1)
    assert xlsx(path, 10000, lambda: False, time.monotonic() + 10).reason == "cell_budget"
    monkeypatch.setattr(extraction_office, "MAX_OFFICE_UNCOMPRESSED", 1)
    with pytest.raises(OverflowError):
        extraction_office.check_office_archive(path)


def test_isolated_runner_has_real_timeout_and_bounded_output(tmp_path: Path) -> None:
    import sys

    runner = ExternalCommandRunner()
    assert runner.run([sys.executable, "-c", "print('synthetic')"], timeout=2).splitlines() == [b"synthetic"]
    assert runner.run([sys.executable, "-c", "import time; time.sleep(3)"], timeout=1) == b""
    assert runner.last_status == "timeout"
    assert (
        runner.run(
            [sys.executable, "-c", "import sys; sys.stdout.write('x' * 17000000)"], timeout=2
        )
        == b""
    )
    assert runner.last_status == "too_large"


def test_immutable_artifact_survives_force_and_snapshot_retry(tmp_path: Path) -> None:
    from papagui_server.adapters.extraction_store import artifact_identity

    store = ExtractionStore(tmp_path / "private" / "artifacts.db")
    original = ExtractionResult(
        text="synthetic original", content_hash="file", version_hash="parser", status="ok"
    )
    original = replace(original, artifact_hash=artifact_identity(original))
    assert store.put(original)
    replacement = replace(original, text="synthetic replacement", artifact_hash="")
    replacement = replace(replacement, artifact_hash=artifact_identity(replacement))
    assert store.put(replacement)
    assert store.get("file", "parser") == replacement
    assert store.read("file", original.artifact_hash) == original
    assert store.read("file", replacement.artifact_hash) == replacement


def test_reader_and_retention_preserve_published_catalog_references(
    tmp_path: Path, monkeypatch
) -> None:
    from papagui_server.adapters.generations import GenerationV2Publisher

    source = tmp_path / "source"
    project = source / "Planung" / "2026" / "Beispiel GmbH"
    project.mkdir(parents=True)
    document = project / "Kontakt.txt"
    document.write_text("synthetic old")
    indexer = SqliteCatalogIndexer(tmp_path / "server")
    _build(indexer, source)
    evidence = SqliteCatalogReader(indexer.active_path).document_evidence(source_id="primary")[0]
    GenerationV2Publisher(tmp_path / "server").publish_index()
    document.write_text("synthetic new")
    for _ in range(4):
        _build(indexer, source)
    later = time.time() + 60 * 86400
    monkeypatch.setattr("papagui_server.adapters.extraction_store.time.time", lambda: later)
    indexer._cleanup_extractions(ServerSettings())
    old = indexer.extraction_store.read(evidence["content_hash"], evidence["extraction_version"])
    assert old and old.text == "synthetic old"
    (indexer.backup_root / "broken.db").write_bytes(b"invalid")
    indexer._cleanup_extractions(ServerSettings())  # Unreadable snapshot: defer cleanup safely.


def test_format_specific_fingerprints_and_reader_source_states(tmp_path: Path) -> None:
    settings = ServerSettings()
    ocr_settings = replace(settings, ocr_max_pages=3)
    assert extraction_fingerprint(settings, "txt") == extraction_fingerprint(ocr_settings, "txt")
    assert extraction_fingerprint(settings, "docx") == extraction_fingerprint(ocr_settings, "docx")
    assert extraction_fingerprint(settings, "pdf") != extraction_fingerprint(ocr_settings, "pdf")
    for extension in ("xls", "xlsx", "doc", "png"):
        assert len(extraction_fingerprint(settings, extension)) == 64
    project = tmp_path / "source" / "Planung" / "2026" / "Beispiel GmbH"
    project.mkdir(parents=True)
    (project / "empty.txt").write_text("")
    indexer = SqliteCatalogIndexer(tmp_path / "server")
    _build(indexer, tmp_path / "source")
    reader = SqliteCatalogReader(indexer.active_path)
    states = reader.source_states(source_id="primary")
    assert states[0]["extraction_status"] == "no_text"
    assert len(states[0]["content_hash"]) == 64
    assert states[0]["source"]["relative_path"].endswith("empty.txt")
    assert reader.source_states(source_id="primary", project_root_ids=[]) == []


def test_partial_transient_cache_retries_and_cancelled_cache_resumes(tmp_path: Path) -> None:
    from papagui_server.domain.document_extraction import ExtractionPage

    store = ExtractionStore(tmp_path / "private" / "artifacts.db")
    result = ExtractionResult(
        status="partial", reason="cancelled", content_hash="file", version_hash="version"
    )
    store.put(result)
    assert store.get("file", "version") is None
    result = replace(result, reason="render_failed", pages=(ExtractionPage(1, "tool_missing"),))
    store.put(result, retry_delay_seconds=0)
    assert store.get("file", "version") is None


def test_pdf_failure_pages_keep_preceding_native_text(tmp_path: Path, monkeypatch) -> None:
    from papagui_server.domain.document_extraction import ExtractionPage

    path = tmp_path / "synthetic.pdf"
    path.write_bytes(b"fixture")
    native = ExtractionResult(
        text="synthetic native " * 20,
        blocks=(ExtractionBlock("synthetic native " * 20, page=1),),
        status="ok",
        pages_total=2,
        pages_processed=2,
        pages=(ExtractionPage(1), ExtractionPage(2, "no_text", width=100, height=200)),
    )
    runner = PageRunner()
    extractor = DocumentTextExtractor(runner)
    monkeypatch.setattr(extractor, "_worker", lambda *_: native)
    result = extractor.extract_document(path, ServerSettings(ocr_enabled=False))
    assert result.text == native.text
    for failing_status in ("tool_missing", "timeout"):

        def fail_render(*_, status=failing_status, **__):
            runner.last_status = status
            return b""

        monkeypatch.setattr(runner, "run", fail_render)
        result = extractor.extract_document(path, ServerSettings(ocr_extension_threshold=100))
        assert result.status == "partial"
        assert result.pages[1].status == failing_status
        assert result.text == native.text
    monkeypatch.setattr(runner, "run", PageRunner().run)
    monkeypatch.setattr(extractor, "_ocr", lambda *_args, **_kwargs: [])
    result = extractor.extract_document(path, ServerSettings(ocr_extension_threshold=100))
    assert result.status == "partial" and result.reason == "ocr_no_text"

    def timeout(*_, **__):
        raise TimeoutError

    monkeypatch.setattr(extractor, "_ocr", timeout)
    result = extractor.extract_document(path, ServerSettings(ocr_extension_threshold=100))
    assert result.status == "partial" and result.reason == "document_time_budget"
    assert result.text == native.text


def test_pdf_two_parties_in_columns_stay_separate_with_coordinates(tmp_path: Path) -> None:
    from reportlab.pdfgen.canvas import Canvas
    from papagui_server.domain.customer_recognition import document_candidates

    path = tmp_path / "synthetic-columns.pdf"
    canvas = Canvas(str(path))
    canvas.setFont("Helvetica", 11)
    left = ["Absender: Eigenes Buero", "Telefon: 030 99999999", "Mail: own@example.org"]
    right = ["Auftraggeber: Beispiel GmbH", "Telefon: 030 12345678", "Mail: kunde@example.org"]
    for number, (own, customer) in enumerate(zip(left, right)):
        canvas.drawString(30, 750 - number * 16, own)
        canvas.drawString(330, 750 - number * 16, customer)
    canvas.save()
    result = read_pdf(path, 10000, 10)
    assert len(result.blocks) == 2
    own, customer = result.blocks
    assert own.bbox[2] < customer.bbox[0]
    assert own.bbox[0] == 30 and customer.bbox[0] == 330
    assert "own@example.org" in own.text and "kunde@example.org" not in own.text
    assert "kunde@example.org" in customer.text and "own@example.org" not in customer.text
    candidates = document_candidates(
        result.text,
        customer_name="Beispiel GmbH",
        own_identities=("Eigenes Buero",),
        blocks=result.blocks,
    )
    assert any(
        candidate.field_name == "email"
        and candidate.value == "kunde@example.org"
        and candidate.quality == "strong"
        for candidate in candidates
    )
    assert not any(
        candidate.field_name == "email"
        and candidate.value == "own@example.org"
        and candidate.quality == "strong"
        for candidate in candidates
    )


def test_worker_memory_budget_is_forwarded_and_changes_cache(tmp_path: Path, monkeypatch) -> None:
    settings = ServerSettings(extraction_memory_mb=256)
    assert extraction_fingerprint(settings, "pdf") != extraction_fingerprint(
        ServerSettings(), "pdf"
    )
    extractor = DocumentTextExtractor()
    commands = []

    def run(command, *, timeout):
        commands.append(command)
        return b'{"status":"no_text"}'

    monkeypatch.setattr(extractor._worker_runner, "run", run)
    extractor._worker(
        "extraction_pdf_worker", tmp_path / "synthetic.pdf", settings, time.monotonic() + 10
    )
    assert commands[0][-1] == "256"


def test_native_character_partial_cannot_turn_into_success_without_ocr(
    tmp_path: Path, monkeypatch
) -> None:
    from papagui_server.domain.document_extraction import ExtractionPage

    extractor = DocumentTextExtractor(PageRunner())
    native = ExtractionResult(
        text="0123456789",
        blocks=(ExtractionBlock("0123456789", page=1),),
        status="partial",
        reason="character_budget",
        pages_total=1,
        pages_processed=1,
        pages=(ExtractionPage(1),),
    )
    monkeypatch.setattr(extractor, "_worker", lambda *_: native)
    result = extractor.extract_document(
        tmp_path / "synthetic.pdf",
        ServerSettings(max_extracted_characters=10, ocr_extension_threshold=5),
    )
    assert result.status == "partial" and result.reason == "character_budget"


def test_installing_ocr_language_invalidates_its_cache(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("TESSDATA_PREFIX", str(tmp_path))
    previous = extraction_fingerprint(ServerSettings(), "png")
    # Only a synthetic file-stat fixture; it is never loaded as a Tesseract model.
    (tmp_path / "deu.traineddata").write_bytes(b"synthetic version marker")
    assert extraction_fingerprint(ServerSettings(), "png") != previous
