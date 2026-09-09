"""Office readers preserving paragraphs, table rows, sheets and display values."""

from __future__ import annotations

from collections.abc import Callable
from datetime import date, datetime
from pathlib import Path
import re
import time
import xml.etree.ElementTree as ET
import zipfile

from papagui_server.domain.document_extraction import (
    ExtractionBlock,
    ExtractionResult,
    bounded_result,
)

MAX_OFFICE_UNCOMPRESSED = 128 * 1024 * 1024
MAX_CELLS = 200_000


def check_office_archive(path: Path) -> None:
    with path.open("rb") as stream:
        if stream.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ValueError("encrypted_office_document")
    with zipfile.ZipFile(path) as bundle:
        if sum(info.file_size for info in bundle.infolist()) > MAX_OFFICE_UNCOMPRESSED:
            raise OverflowError("archive_budget")


def docx(
    path: Path, maximum: int, cancelled: Callable[[], bool], deadline: float
) -> ExtractionResult:
    check_office_archive(path)
    blocks: list[ExtractionBlock] = []
    length = 0
    with zipfile.ZipFile(path) as bundle:
        names = bundle.namelist()
        sections = sorted(name for name in names if re.fullmatch(r"word/header\d+\.xml", name))
        sections += ["word/document.xml"]
        sections += sorted(name for name in names if re.fullmatch(r"word/footer\d+\.xml", name))
        for name in sections:
            if name not in names:
                continue
            root = ET.fromstring(bundle.read(name))
            section = Path(name).stem
            # Runs in one paragraph are joined without introducing false line breaks.
            parents = {child: parent for parent in root.iter() for child in parent}
            for node in root.iter():
                tag = node.tag.rsplit("}", 1)[-1]
                if tag not in {"p", "tr"}:
                    continue
                if tag == "p":
                    ancestor = parents.get(node)
                    in_row = False
                    while ancestor is not None:
                        if ancestor.tag.rsplit("}", 1)[-1] == "tr":
                            in_row = True
                            break
                        ancestor = parents.get(ancestor)
                    if in_row:
                        continue
                if cancelled() or time.monotonic() > deadline:
                    return bounded_result(
                        blocks,
                        maximum,
                        status="partial",
                        reason="cancelled" if cancelled() else "time_budget",
                    )
                if tag == "tr":
                    value = "\t".join(
                        _paragraph_text(cell)
                        for cell in node
                        if cell.tag.rsplit("}", 1)[-1] == "tc"
                    )
                else:
                    value = _paragraph_text(node)
                if value.strip():
                    blocks.append(
                        ExtractionBlock(
                            value,
                            kind="table_row" if tag == "tr" else "paragraph",
                            section=section,
                            method="docx",
                        )
                    )
                    length += len(value) + 1
                    if length > maximum:
                        return bounded_result(blocks, maximum)
    return bounded_result(blocks, maximum)


def _paragraph_text(node: ET.Element) -> str:
    parts: list[str] = []
    # The unnamespaced branch permits small synthetic XML fixtures.
    if "}" not in node.tag:
        return "".join(node.itertext())
    for child in node.iter():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "t":
            parts.append(child.text or "")
        elif tag == "tab":
            parts.append("\t")
        elif tag in {"br", "cr"}:
            parts.append("\n")
        elif tag == "p" and child is not node and parts:
            parts.append("\n")
    return "".join(parts).strip()


def display_value(value: object, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        # Preserve displayed identifiers, postal codes and phone numbers, but do
        # not manufacture precision from accounting, date or scientific formats.
        mask = number_format.split(";", 1)[0]
        mask = re.sub(r'"([^"]*)"', r"\1", mask)
        mask = mask.replace("\\", "")
        if (
            float(value).is_integer()
            and re.fullmatch(r"[0 ()+./-]+", mask)
            and "0" in mask
            and "." not in mask
        ):
            digits = str(abs(int(value))).zfill(mask.count("0"))
            if len(digits) <= mask.count("0"):
                iterator = iter(digits)
                return ("-" if value < 0 else "") + "".join(
                    next(iterator) if char == "0" else char for char in mask
                )
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def xlsx(
    path: Path, maximum: int, cancelled: Callable[[], bool], deadline: float
) -> ExtractionResult:
    import openpyxl

    check_office_archive(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = None
    blocks: list[ExtractionBlock] = []
    length = cells = 0
    missing_formula = False
    try:
        # A second streaming workbook distinguishes a blank cell from a formula
        # for which Excel has not saved a cached result. No formula is executed.
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        for sheet, formula_sheet in zip(workbook.worksheets, formulas.worksheets):
            for row, formula_row in zip(sheet.iter_rows(), formula_sheet.iter_rows()):
                cells += len(row)
                if cells > MAX_CELLS or cancelled() or time.monotonic() > deadline:
                    reason = (
                        "cell_budget"
                        if cells > MAX_CELLS
                        else "cancelled"
                        if cancelled()
                        else "time_budget"
                    )
                    return bounded_result(blocks, maximum, status="partial", reason=reason)
                for cell, formula in zip(row, formula_row):
                    if formula.data_type == "f" and cell.value is None:
                        missing_formula = True
                        continue
                    value = display_value(cell.value, cell.number_format)
                    if value:
                        blocks.append(
                            ExtractionBlock(
                                value,
                                kind="cell",
                                sheet=sheet.title,
                                cell=cell.coordinate,
                                method="openpyxl",
                            )
                        )
                        length += len(value) + 1
                        if length > maximum:
                            return bounded_result(blocks, maximum)
    finally:
        workbook.close()
        if formulas is not None:
            formulas.close()
    return bounded_result(
        blocks,
        maximum,
        status="partial" if missing_formula else "ok",
        reason="formula_result_missing" if missing_formula else "",
    )


def xls(
    path: Path, maximum: int, cancelled: Callable[[], bool], deadline: float
) -> ExtractionResult:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
    blocks: list[ExtractionBlock] = []
    length = cells = 0
    try:
        for sheet in workbook.sheets():
            for row in range(sheet.nrows):
                for column in range(sheet.ncols):
                    cells += 1
                    if cells > MAX_CELLS or cancelled() or time.monotonic() > deadline:
                        reason = (
                            "cell_budget"
                            if cells > MAX_CELLS
                            else "cancelled"
                            if cancelled()
                            else "time_budget"
                        )
                        return bounded_result(blocks, maximum, status="partial", reason=reason)
                    cell = sheet.cell(row, column)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(int(value), "#ERROR!")
                    fmt = workbook.format_map.get(workbook.xf_list[cell.xf_index].format_key)
                    text = display_value(value, fmt.format_str if fmt else "")
                    if text:
                        col = column + 1
                        letters = ""
                        while col:
                            col, remainder = divmod(col - 1, 26)
                            letters = chr(65 + remainder) + letters
                        blocks.append(
                            ExtractionBlock(
                                text,
                                kind="cell",
                                sheet=sheet.name,
                                cell=f"{letters}{row + 1}",
                                method="xlrd",
                            )
                        )
                        length += len(text) + 1
                        if length > maximum:
                            return bounded_result(blocks, maximum)
    finally:
        workbook.release_resources()
    return bounded_result(blocks, maximum)
