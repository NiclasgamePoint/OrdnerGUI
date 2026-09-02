from __future__ import annotations

import os
from pathlib import Path
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl
from pypdf import PdfWriter
from PySide6.QtCore import Qt
from PySide6.QtGui import QPixmap
from PySide6.QtWidgets import QApplication
import pytest

from papagui_contracts import CatalogFacets, CatalogFolder, CatalogProjectRoot, SourcePath
from papagui_client.application.models import (
    CatalogFolderDetails,
    CatalogFolderHit,
    CatalogHit,
    CatalogProjectRootHit,
    CatalogRecord,
    GlobalSearchHit,
    GlobalSearchKind,
    GlobalSearchPage,
    GlobalSearchRecord,
)
from papagui_client.gui.viewers import CatalogBrowserWidget, FileViewerWidget
from papagui_client.viewers.models import ConversionOutcome


@pytest.fixture(scope="module")
def application():
    return QApplication.instance() or QApplication([])


def _process_until(application, predicate, timeout: float = 3) -> None:
    deadline = time.monotonic() + timeout
    while not predicate() and time.monotonic() < deadline:
        application.processEvents()
        time.sleep(0.01)
    application.processEvents()
    assert predicate()


def _pdf(path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    with path.open("wb") as handle:
        writer.write(handle)


def test_file_viewer_routes_pdf_text_spreadsheet_image_and_unknown(application, tmp_path):
    pdf = tmp_path / "document.pdf"
    _pdf(pdf)
    text = tmp_path / "notes.txt"
    text.write_text("eins zwei drei", encoding="utf-8")
    spreadsheet = tmp_path / "table.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Daten"
    workbook.active["B2"] = "Treffer"
    workbook.save(spreadsheet)
    image = tmp_path / "picture.png"
    pixmap = QPixmap(8, 4)
    pixmap.fill(Qt.GlobalColor.blue)
    assert pixmap.save(str(image))
    unknown = tmp_path / "archive.bin"
    unknown.write_bytes(b"unknown")

    viewer = FileViewerWidget(preview_cache_root=tmp_path / "cache")
    assert viewer.open_file(pdf)
    assert viewer.pdf_viewer.document.pageCount() == 1
    viewer.open_file(spreadsheet)
    assert viewer.spreadsheet_viewer.table.item(1, 1).text() == "Treffer"
    viewer.open_file(text)
    viewer.text_viewer.search_input.setText("zwei")
    viewer.text_viewer.find_next()
    assert viewer.text_viewer.editor.textCursor().hasSelection()
    viewer.open_file(image)
    assert viewer.stack.currentWidget() is viewer.image_viewer
    viewer.open_file(unknown)
    assert "keine interne Vorschau" in viewer.text_viewer.editor.toPlainText()
    viewer.shutdown()
    viewer.close()


class _FailingWordConverter:
    def convert_word_to_pdf(self, *_args, **_kwargs):
        raise RuntimeError("kein Konverter")

    def cleanup(self):
        return None


class _SuccessfulWordConverter:
    def __init__(self, pdf: Path):
        self.pdf = pdf

    def convert_word_to_pdf(self, *_args, **_kwargs):
        return ConversionOutcome("word_to_pdf", "Test", True, path=self.pdf)

    def cleanup(self):
        return None


def test_docx_preview_falls_back_to_text_asynchronously(application, tmp_path):
    docx = pytest.importorskip("docx")
    source = tmp_path / "letter.docx"
    document = docx.Document()
    document.add_paragraph("Fallback Inhalt")
    document.save(source)
    viewer = FileViewerWidget(converter_factory=_FailingWordConverter)

    viewer.open_file(source)
    assert viewer.is_converting
    _process_until(
        application,
        lambda: viewer.stack.currentWidget() is viewer.text_viewer and not viewer.is_converting,
    )
    assert "kein Konverter" in viewer.text_viewer.editor.toPlainText()
    assert "Fallback Inhalt" in viewer.text_viewer.editor.toPlainText()
    viewer.shutdown()
    viewer.close()


def test_word_preview_uses_background_conversion(application, tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"source")
    rendered = tmp_path / "rendered.pdf"
    _pdf(rendered)
    viewer = FileViewerWidget(
        converter_factory=lambda: _SuccessfulWordConverter(rendered)
    )

    viewer.open_file(source)
    _process_until(
        application,
        lambda: viewer.stack.currentWidget() is viewer.pdf_viewer and not viewer.is_converting,
    )
    assert viewer.pdf_viewer.document.pageCount() == 1
    assert "Test" in viewer.conversion_meta_label.text()
    viewer.shutdown()
    viewer.close()


class _CatalogSearch:
    def __init__(self, hits):
        self.hits = hits
        self.query = None

    def search(self, query, **_filters):
        self.query = query
        return self.hits


def test_catalog_browser_filters_catalog_hits_and_previews_local_path(application, tmp_path):
    source = tmp_path / "angebot.txt"
    source.write_text("Lokale Vorschau", encoding="utf-8")
    hit = CatalogHit(
        CatalogRecord(
            document_key="doc-1",
            source_id="archive",
            relative_path="2026/angebot.txt",
            filename="angebot.txt",
            file_type="txt",
            customer_name="Muster GmbH",
            project_name="Projekt A",
            year="2026",
        ),
        source,
    )
    search = _CatalogSearch([hit])
    browser = CatalogBrowserWidget(search)
    browser.query_input.setText("Angebot")
    browser.customer_input.setText("Muster GmbH")
    browser.year_input.setText("2026")
    browser.file_type_input.setCurrentText("txt")

    browser.run_search()
    application.processEvents()

    assert browser.results.topLevelItemCount() == 1
    assert search.query.text == "Angebot"
    assert search.query.customer_name == "Muster GmbH"
    assert search.query.year == "2026"
    assert browser.selected_hit() == hit
    assert "Lokale Vorschau" in browser.viewer.text_viewer.editor.toPlainText()
    assert browser.status_label.text() == "1 Treffer"
    browser.shutdown()
    browser.close()


def test_catalog_browser_handles_unavailable_file_and_search_error(application, tmp_path):
    missing = tmp_path / "missing.pdf"
    hit = CatalogHit(
        CatalogRecord("missing", "archive", "missing.pdf", "missing.pdf", "pdf"),
        missing,
    )
    browser = CatalogBrowserWidget(_CatalogSearch([hit]))
    browser.run_search()
    application.processEvents()
    assert "nicht erreichbar" in browser.viewer.text_viewer.editor.toPlainText()

    class BrokenSearch:
        def search(self, *_args, **_kwargs):
            raise OSError("offline")

    failed = CatalogBrowserWidget(BrokenSearch())
    failed.run_search()
    assert "offline" in failed.status_label.text()
    failed.shutdown()
    failed.close()
    browser.shutdown()
    browser.close()


class _GlobalSearch:
    def __init__(self, root, document):
        source = SourcePath("archive", "2026/Muster")
        self.root = CatalogProjectRoot(1, source, "Beratung", 2026, "Muster", "Muster")
        self.folder_value = CatalogFolder(1, source, "Muster", file_count=1)
        self.child = CatalogFolder(
            2,
            SourcePath("archive", "2026/Muster/Pläne"),
            "Pläne",
            parent_id=1,
        )
        self.root_hit = CatalogProjectRootHit(self.root, root)
        self.folder_hit = CatalogFolderHit(self.folder_value, root)
        self.child_hit = CatalogFolderHit(self.child, root / "Pläne")
        self.document_hit = CatalogHit(
            CatalogRecord(
                "doc",
                "archive",
                "2026/Muster/angebot.txt",
                "angebot.txt",
                "txt",
            ),
            document,
        )
        self.offsets = []

    def facets(self, _source=None):
        return CatalogFacets(("Beratung",), ("2026",), ("txt",))

    def history(self):
        return ("Muster",)

    def project_roots(self):
        return (self.root_hit,)

    def folders(self):
        return (self.folder_hit, self.child_hit)

    def folder(self, _source_id, _relative_path):
        return CatalogFolderDetails(
            self.folder_hit, (self.child_hit,), (self.document_hit,)
        )

    def global_search(self, query):
        self.offsets.append(query.offset)
        all_items = (
            GlobalSearchHit(
                GlobalSearchRecord(
                    GlobalSearchKind.CUSTOMER,
                    "customer:7",
                    "Muster GmbH",
                    customer_id=7,
                )
            ),
            GlobalSearchHit(
                GlobalSearchRecord(
                    GlobalSearchKind.FOLDER,
                    "folder:archive:1",
                    "Muster",
                    source_id="archive",
                    relative_path="2026/Muster",
                ),
                self.folder_hit.local_path,
            ),
        )
        return GlobalSearchPage(all_items, 202, query.limit, query.offset)


def test_catalog_browser_global_search_navigation_activation_and_paging(
    application, tmp_path
):
    document = tmp_path / "angebot.txt"
    document.write_text("Globaler Inhalt", encoding="utf-8")
    search = _GlobalSearch(tmp_path, document)
    browser = CatalogBrowserWidget(search)
    assert browser.domain_filter.findData("Beratung") >= 0
    assert browser.navigation.topLevelItemCount() == 1
    assert browser._history_model.stringList() == ["Muster"]

    activated = []
    browser.resultActivated.connect(activated.append)
    browser.query_input.setText("Muster")
    browser.kind_filter.setCurrentIndex(4)
    browser.domain_filter.setCurrentIndex(1)
    browser.file_type_input.setCurrentText(".txt")
    browser.run_search()
    assert browser.results.topLevelItemCount() == 2
    assert browser.next_button.isEnabled()
    assert "1–2" in browser.page_label.text()
    browser.next_page()
    assert search.offsets[-1] == 100
    browser.previous_page()
    assert search.offsets[-1] == 0

    browser.results.setCurrentItem(browser.results.topLevelItem(0))
    browser._emit_activated(browser.results.topLevelItem(0), 0)
    assert activated[-1].record.kind is GlobalSearchKind.CUSTOMER
    browser.results.setCurrentItem(browser.results.topLevelItem(1))
    browser._emit_activated(browser.results.topLevelItem(1), 0)
    assert "Unterordner" in browser.status_label.text()

    root_item = browser.navigation.topLevelItem(0)
    browser._open_navigation_item(root_item, 0)
    assert browser.results.topLevelItemCount() == 2
    browser.results.setCurrentItem(browser.results.topLevelItem(1))
    browser.preview_selected()
    assert "Globaler Inhalt" in browser.viewer.text_viewer.editor.toPlainText()
    browser.close()
