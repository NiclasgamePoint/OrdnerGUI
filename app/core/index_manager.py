import sqlite3
from pathlib import Path
from datetime import datetime
import os
from typing import Callable, List, Dict, Optional
import subprocess
import shutil
import time
import re
import sys
import json
import hashlib
import importlib.util
import logging
from concurrent.futures import ThreadPoolExecutor, as_completed
from tempfile import TemporaryDirectory

import openpyxl
import xlrd
from docx import Document
from PyPDF2 import PdfReader
from app.core.config import IndexOptions
from app.core.search_models import SearchFilters, SearchPage


logger = logging.getLogger(__name__)


class IndexManager:
    SEARCH_LABEL_SQL = "COALESCE(NULLIF(project_name, ''), NULLIF(customer_name, ''), filename)"
    BINARY_CONTENT_TYPES = {"pdf", "doc", "docx", "xls", "xlsx"}
    TEXT_CONTENT_TYPES = {
        "txt", "csv", "md", "log", "json", "xml", "yaml", "yml", "ini"
    }
    CONTENT_INDEX_TYPES = BINARY_CONTENT_TYPES | TEXT_CONTENT_TYPES
    SCHEMA_VERSION = "3"
    EXTRACTOR_VERSION = "3"
    APP_VERSION = "0.2"

    def __init__(
        self,
        db_path: Path,
        initialize: bool = True,
        options: Optional[IndexOptions] = None,
    ):
        self.db_path = db_path
        self.conn = None
        self.options = options or IndexOptions()
        self._ocr_language = None
        if initialize:
            self.init_db()
        else:
            self.conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
            self.conn.row_factory = sqlite3.Row
    
    def init_db(self):
        self.conn = sqlite3.connect(str(self.db_path))
        self.conn.row_factory = sqlite3.Row
        cursor = self.conn.cursor()
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS files (
                id INTEGER PRIMARY KEY,
                path TEXT UNIQUE NOT NULL,
                filename TEXT NOT NULL,
                file_size INTEGER,
                file_type TEXT,
                created_date TIMESTAMP,
                modified_date TIMESTAMP,
                year TEXT,
                service_type TEXT,
                customer_name TEXT,
                subfolder TEXT,
                full_text_indexed BOOLEAN DEFAULT 0,
                domain_folder TEXT,
                time_bucket TEXT,
                project_name TEXT,
                relative_dir TEXT,
                index_root TEXT
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS indexed_roots (
                root_path TEXT PRIMARY KEY,
                last_indexed TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS index_metadata (
                key TEXT PRIMARY KEY,
                value TEXT
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS folders (
                path TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                relative_path TEXT,
                parent_path TEXT,
                index_root TEXT NOT NULL
            )
        """)

        cursor.execute("""
            CREATE VIRTUAL TABLE IF NOT EXISTS file_content_fts USING fts5(
                path UNINDEXED,
                content,
                tokenize = 'unicode61 remove_diacritics 2'
            )
        """)

        self._ensure_column(cursor, "files", "domain_folder", "TEXT")
        self._ensure_column(cursor, "files", "time_bucket", "TEXT")
        self._ensure_column(cursor, "files", "project_name", "TEXT")
        self._ensure_column(cursor, "files", "relative_dir", "TEXT")
        self._ensure_column(cursor, "files", "index_root", "TEXT")
        self._ensure_column(cursor, "files", "folder_path", "TEXT")
        self._ensure_column(cursor, "files", "modified_ns", "INTEGER")
        self._ensure_column(cursor, "files", "content_hash", "TEXT")
        self._ensure_column(cursor, "files", "content_status", "TEXT")
        self._ensure_column(cursor, "files", "content_error", "TEXT")
        self._ensure_column(cursor, "files", "extractor_version", "TEXT")
        
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_customer ON files(customer_name)
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_service_type ON files(service_type)
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_year ON files(year)
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_filename ON files(filename)
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_domain_folder ON files(domain_folder)
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_time_bucket ON files(time_bucket)
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_project_name ON files(project_name)
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_index_root ON files(index_root)
        """)
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_folder_path ON files(folder_path)")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_folders_name ON folders(name)")
        
        self.conn.commit()

    def _ensure_column(self, cursor: sqlite3.Cursor, table: str, column: str, col_type: str):
        cursor.execute(f"PRAGMA table_info({table})")
        columns = {row[1] for row in cursor.fetchall()}
        if column not in columns:
            cursor.execute(f"ALTER TABLE {table} ADD COLUMN {column} {col_type}")

    def _is_year_bucket(self, value: Optional[str]) -> bool:
        if not value:
            return False
        return len(value) == 4 and value.isdigit()

    def _derive_project_from_filename(self, filename: str) -> str:
        stem = Path(filename).stem.strip()
        separators = [" vom ", " - ", ","]
        for sep in separators:
            if sep in stem:
                left = stem.split(sep, 1)[0].strip()
                if left:
                    return left
        return stem

    def has_index_for_root(self, base_path: Path) -> bool:
        cursor = self.conn.cursor()
        root = str(base_path)
        cursor.execute("SELECT 1 FROM indexed_roots WHERE root_path = ? LIMIT 1", (root,))
        return cursor.fetchone() is not None

    def content_index_needs_rebuild(self, base_path: Path) -> bool:
        cursor = self.conn.cursor()
        content_types = self.CONTENT_INDEX_TYPES & self.options.indexed_content_types
        if not content_types:
            return False
        placeholders = ", ".join("?" for _ in content_types)
        cursor.execute(
            f"""
            SELECT 1 FROM files
            WHERE index_root = ?
              AND file_type IN ({placeholders})
              AND COALESCE(full_text_indexed, 0) = 0
            LIMIT 1
            """,
            (str(base_path), *sorted(content_types)),
        )
        return cursor.fetchone() is not None

    def get_metadata(self, key: str, default: str = "") -> str:
        row = self.conn.execute(
            "SELECT value FROM index_metadata WHERE key = ?", (key,)
        ).fetchone()
        return str(row[0]) if row is not None else default

    def set_metadata(self, key: str, value: str):
        self.conn.execute(
            "INSERT OR REPLACE INTO index_metadata (key, value) VALUES (?, ?)",
            (key, str(value)),
        )

    def index_is_current(self, base_path: Path) -> bool:
        return (
            self.has_index_for_root(base_path)
            and self.get_metadata("schema_version") == self.SCHEMA_VERSION
            and self.get_metadata("extractor_version") == self.EXTRACTOR_VERSION
            and self.get_metadata("options_fingerprint") == self.options.fingerprint()
            and not self.content_index_needs_rebuild(base_path)
        )
    
    def index_directory(
        self,
        base_path: Path,
        replace_existing: bool = False,
        should_cancel: Optional[Callable[[], bool]] = None,
    ):
        return self.synchronize_directory(
            base_path,
            full_rebuild=replace_existing,
            should_cancel=should_cancel,
        )

    def synchronize_directory(
        self,
        base_path: Path,
        full_rebuild: bool = False,
        should_cancel: Optional[Callable[[], bool]] = None,
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> int:
        """Incrementally synchronize a source into this (usually staged) index."""
        started_at = time.monotonic()
        base_path = base_path.resolve()
        logger.info("Indexabgleich gestartet: root=%s full_rebuild=%s", base_path, full_rebuild)
        if not base_path.exists() or not base_path.is_dir():
            raise FileNotFoundError(f"Index-Pfad existiert nicht: {base_path}")

        cursor = self.conn.cursor()
        existing_root = self.get_metadata("index_root")
        changed_count = 1 if full_rebuild else 0
        if full_rebuild or (existing_root and existing_root != str(base_path)):
            cursor.execute("DELETE FROM file_content_fts")
            cursor.execute("DELETE FROM files")
            cursor.execute("DELETE FROM folders")
            cursor.execute("DELETE FROM indexed_roots")

        cursor.execute("CREATE TEMP TABLE IF NOT EXISTS seen_files (path TEXT PRIMARY KEY)")
        cursor.execute("DELETE FROM seen_files")
        cursor.execute("CREATE TEMP TABLE IF NOT EXISTS seen_folders (path TEXT PRIMARY KEY)")
        cursor.execute("DELETE FROM seen_folders")

        processed_count = 0
        excluded = self.options.excluded_folder_names
        for current_root, directory_names, file_names in os.walk(base_path):
            if should_cancel is not None and should_cancel():
                raise InterruptedError("Indexierung wurde abgebrochen")

            directory_names[:] = [
                name for name in directory_names if name.casefold() not in excluded
            ]
            current_path = Path(current_root)
            relative = current_path.relative_to(base_path)
            relative_text = "" if relative == Path(".") else str(relative)
            parent_path = str(current_path.parent) if current_path != base_path else ""
            folder_exists = cursor.execute(
                "SELECT 1 FROM folders WHERE path = ?", (str(current_path),)
            ).fetchone()
            cursor.execute(
                """
                INSERT OR REPLACE INTO folders
                (path, name, relative_path, parent_path, index_root)
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    str(current_path),
                    current_path.name,
                    relative_text,
                    parent_path,
                    str(base_path),
                ),
            )
            cursor.execute("INSERT OR IGNORE INTO seen_folders(path) VALUES (?)", (str(current_path),))
            if folder_exists is None:
                changed_count += 1

            for filename in file_names:
                if should_cancel is not None and should_cancel():
                    raise InterruptedError("Indexierung wurde abgebrochen")
                filepath = current_path / filename
                try:
                    stat = filepath.stat()
                except OSError as exc:
                    logger.warning("Datei übersprungen: path=%s error=%s", filepath, exc)
                    continue

                path_text = str(filepath)
                cursor.execute("INSERT OR IGNORE INTO seen_files(path) VALUES (?)", (path_text,))
                existing = cursor.execute(
                    """
                    SELECT file_size, modified_ns, extractor_version, content_error
                    FROM files WHERE path = ?
                    """,
                    (path_text,),
                ).fetchone()
                unchanged = (
                    not full_rebuild
                    and existing is not None
                    and existing["file_size"] == stat.st_size
                    and existing["modified_ns"] == stat.st_mtime_ns
                    and (
                        filepath.suffix.lower().lstrip(".")
                        not in (self.CONTENT_INDEX_TYPES & self.options.indexed_content_types)
                        or existing["extractor_version"] == self.EXTRACTOR_VERSION
                    )
                    and not (
                        "PyCryptodome is required" in (existing["content_error"] or "")
                        and importlib.util.find_spec("Crypto") is not None
                    )
                )
                if not unchanged:
                    self._index_file(filepath, base_path, cursor, stat=stat)
                    changed_count += 1

                processed_count += 1
                if progress_callback is not None:
                    progress_callback(processed_count, path_text)
                if processed_count % 250 == 0:
                    self.conn.commit()

        cursor.execute(
            "DELETE FROM file_content_fts WHERE path IN "
            "(SELECT path FROM files WHERE path NOT IN (SELECT path FROM seen_files))"
        )
        cursor.execute("DELETE FROM files WHERE path NOT IN (SELECT path FROM seen_files)")
        changed_count += max(cursor.rowcount, 0)
        cursor.execute("DELETE FROM folders WHERE path NOT IN (SELECT path FROM seen_folders)")
        changed_count += max(cursor.rowcount, 0)
        changed_count += self._normalize_content_statuses()
        cursor.execute("DELETE FROM indexed_roots")
        cursor.execute(
            "INSERT INTO indexed_roots(root_path, last_indexed) VALUES (?, ?)",
            (str(base_path), datetime.now().isoformat()),
        )
        self.set_metadata("schema_version", self.SCHEMA_VERSION)
        self.set_metadata("extractor_version", self.EXTRACTOR_VERSION)
        self.set_metadata("app_version", self.APP_VERSION)
        self.set_metadata("options_fingerprint", self.options.fingerprint())
        self.set_metadata("index_root", str(base_path))
        self.set_metadata("built_at", datetime.now().isoformat())
        self.set_metadata("build_mode", "full" if full_rebuild else "incremental")
        self.set_metadata("file_count", str(processed_count))
        self.set_metadata("changed_count", str(changed_count))
        self.set_metadata("duration_seconds", f"{time.monotonic() - started_at:.3f}")
        self.conn.commit()
        self.last_change_count = changed_count
        logger.info(
            "Indexabgleich abgeschlossen: files=%s changes=%s duration=%.3fs",
            processed_count,
            changed_count,
            time.monotonic() - started_at,
        )
        return processed_count
    
    def _index_file(
        self,
        filepath: Path,
        base_path: Path,
        cursor: sqlite3.Cursor,
        stat=None,
    ):
        try:
            stat = stat or filepath.stat()
            rel_path = str(filepath.relative_to(base_path))
            parts = rel_path.split(os.sep)
            
            domain_folder = parts[0] if len(parts) > 0 else None
            time_bucket = parts[1] if len(parts) > 1 else None

            project_name = None
            subfolder = None
            if len(parts) >= 4:
                project_name = parts[2]
            elif len(parts) == 3:
                project_name = self._derive_project_from_filename(filepath.name)

            if len(parts) >= 5:
                subfolder = parts[3]

            year = time_bucket if self._is_year_bucket(time_bucket) else None
            service_type = domain_folder
            customer_name = project_name or self._derive_project_from_filename(filepath.name)
            
            file_type = filepath.suffix.lower().lstrip('.')
            relative_dir = str(filepath.parent.relative_to(base_path))
            
            cursor.execute("""
                INSERT OR REPLACE INTO files 
                (path, filename, file_size, file_type, created_date, modified_date, 
                 year, service_type, customer_name, subfolder,
                 domain_folder, time_bucket, project_name, relative_dir, index_root,
                 full_text_indexed, folder_path, modified_ns, content_hash,
                 content_status, content_error, extractor_version)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                str(filepath),
                filepath.name,
                stat.st_size,
                file_type,
                datetime.fromtimestamp(stat.st_ctime).isoformat(),
                datetime.fromtimestamp(stat.st_mtime).isoformat(),
                year,
                service_type,
                customer_name,
                subfolder,
                domain_folder,
                time_bucket,
                project_name,
                relative_dir,
                str(base_path),
                0,
                str(filepath.parent),
                stat.st_mtime_ns,
                "",
                "not_applicable",
                "",
                "",
            ))

            if file_type in (self.CONTENT_INDEX_TYPES & self.options.indexed_content_types):
                cursor.execute("DELETE FROM file_content_fts WHERE path = ?", (str(filepath),))
                max_bytes = self.options.max_file_size_mb * 1024 * 1024
                if stat.st_size > max_bytes:
                    content = ""
                    status = "skipped_large"
                    error = f"Datei größer als {self.options.max_file_size_mb} MB"
                else:
                    content, status, error = self._extract_document_with_status(
                        filepath, file_type
                    )
                if content:
                    cursor.execute(
                        "INSERT INTO file_content_fts (path, content) VALUES (?, ?)",
                        (str(filepath), content),
                    )
                cursor.execute(
                    """
                    UPDATE files
                    SET full_text_indexed = 1,
                        content_hash = ?,
                        content_status = ?,
                        content_error = ?,
                        extractor_version = ?
                    WHERE path = ?
                    """,
                    (
                        self._hash_file(filepath) if stat.st_size <= max_bytes else "",
                        status,
                        error,
                        self.EXTRACTOR_VERSION,
                        str(filepath),
                    ),
                )
        except Exception:
            logger.exception("Fehler beim Indexieren: path=%s", filepath)

    def _hash_file(self, filepath: Path) -> str:
        digest = hashlib.sha256()
        with filepath.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def _limit_text(self, text: str) -> str:
        return text[:self.options.max_extracted_characters]

    def _extract_document_with_status(
        self, filepath: Path, file_type: str
    ) -> tuple[str, str, str]:
        try:
            content = self._extract_document_text(filepath, file_type)
            return content, ("success" if content.strip() else "empty"), ""
        except TimeoutError as exc:
            return "", "timeout", str(exc)
        except Exception as exc:
            message = str(exc)
            if "encrypted" in message.casefold() or "password" in message.casefold():
                return "", "encrypted", message
            return "", "error", message

    def _normalize_content_statuses(self) -> int:
        """Migrate older generic errors into actionable diagnostic categories."""
        changed = 0
        cursor = self.conn.cursor()
        cursor.execute(
            """
            UPDATE files SET content_status = 'encrypted'
            WHERE content_status = 'error'
              AND (LOWER(content_error) LIKE '%encrypted%'
                   OR LOWER(content_error) LIKE '%password%')
            """
        )
        changed += max(cursor.rowcount, 0)
        cursor.execute(
            """
            UPDATE files SET content_status = 'timeout'
            WHERE content_status = 'error'
              AND (LOWER(content_error) LIKE '%timed out%'
                   OR LOWER(content_error) LIKE '%futures unfinished%'
                   OR LOWER(content_error) LIKE '%ocr-zeitlimit%')
            """
        )
        changed += max(cursor.rowcount, 0)
        return changed

    def _extract_document_text(self, filepath: Path, file_type: str) -> str:
        if file_type in self.TEXT_CONTENT_TYPES:
            return self._limit_text(filepath.read_text(encoding="utf-8", errors="replace"))
        if file_type == "pdf":
            return self._extract_pdf_text(filepath)
        if file_type == "docx":
            return self._extract_docx_text(filepath)
        if file_type == "doc":
            return self._extract_doc_text(filepath)
        if file_type == "xlsx":
            return self._extract_xlsx_text(filepath)
        if file_type == "xls":
            return self._extract_xls_text(filepath)
        return ""

    def _extract_pdf_text(self, filepath: Path) -> str:
        reader = PdfReader(str(filepath))
        parts = []
        length = 0
        for page in reader.pages:
            text = page.extract_text() or ""
            if text:
                parts.append(text)
                length += len(text)
            if length >= self.options.max_extracted_characters:
                break
        extracted = self._limit_text("\n".join(parts))
        if extracted.strip() or not self.options.ocr_enabled:
            return extracted
        return self._ocr_pdf(filepath)

    def _extract_docx_text(self, filepath: Path) -> str:
        document = Document(str(filepath))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        length = sum(map(len, parts))
        for table in document.tables:
            for row in table.rows:
                values = [cell.text for cell in row.cells if cell.text]
                if values:
                    line = "\t".join(values)
                    parts.append(line)
                    length += len(line)
                if length >= self.options.max_extracted_characters:
                    return self._limit_text("\n".join(parts))
        return self._limit_text("\n".join(parts))

    def _extract_xlsx_text(self, filepath: Path) -> str:
        workbook = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
        try:
            parts = []
            length = 0
            for worksheet in workbook.worksheets:
                parts.append(worksheet.title)
                for row in worksheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        line = "\t".join(values)
                        parts.append(line)
                        length += len(line)
                    if length >= self.options.max_extracted_characters:
                        return self._limit_text("\n".join(parts))
            return self._limit_text("\n".join(parts))
        finally:
            workbook.close()

    def _extract_xls_text(self, filepath: Path) -> str:
        workbook = xlrd.open_workbook(str(filepath), on_demand=True)
        try:
            parts = []
            length = 0
            for worksheet in workbook.sheets():
                parts.append(worksheet.name)
                for row_index in range(worksheet.nrows):
                    values = [
                        str(worksheet.cell_value(row_index, column_index))
                        for column_index in range(worksheet.ncols)
                        if worksheet.cell_value(row_index, column_index) != ""
                    ]
                    if values:
                        line = "\t".join(values)
                        parts.append(line)
                        length += len(line)
                    if length >= self.options.max_extracted_characters:
                        return self._limit_text("\n".join(parts))
            return self._limit_text("\n".join(parts))
        finally:
            workbook.release_resources()

    def _extract_doc_text(self, filepath: Path) -> str:
        for executable_name in ("catdoc", "antiword"):
            executable = shutil.which(executable_name)
            if executable:
                result = subprocess.run(
                    [executable, str(filepath)],
                    capture_output=True,
                    text=True,
                    errors="replace",
                    timeout=30,
                )
                if result.returncode == 0 and result.stdout.strip():
                    return self._limit_text(result.stdout)

        libreoffice = self._find_libreoffice()
        if libreoffice is None:
            return ""

        with TemporaryDirectory(prefix="papagui-doc-index-") as temp_dir:
            output_dir = Path(temp_dir)
            profile_dir = output_dir / "profile"
            runtime_dir = output_dir / "runtime"
            config_dir = output_dir / "config"
            cache_dir = output_dir / "cache"
            profile_dir.mkdir()
            runtime_dir.mkdir(mode=0o700)
            config_dir.mkdir()
            cache_dir.mkdir()
            environment = os.environ.copy()
            environment.update({
                "XDG_RUNTIME_DIR": str(runtime_dir),
                "XDG_CONFIG_HOME": str(config_dir),
                "XDG_CACHE_HOME": str(cache_dir),
                "SAL_USE_VCLPLUGIN": "svp",
            })
            result = subprocess.run(
                [
                    str(libreoffice),
                    "--headless",
                    f"-env:UserInstallation={profile_dir.as_uri()}",
                    "--convert-to",
                    "txt:Text",
                    "--outdir",
                    str(output_dir),
                    str(filepath),
                ],
                capture_output=True,
                text=True,
                timeout=45,
                env=environment,
            )
            converted = next(output_dir.glob("*.txt"), None)
            if result.returncode == 0 and converted is not None:
                return self._limit_text(converted.read_text(encoding="utf-8", errors="replace"))
        return ""

    def _ocr_pdf(self, filepath: Path) -> str:
        pdftoppm = shutil.which("pdftoppm")
        tesseract = self._find_tesseract()
        if pdftoppm is None or tesseract is None:
            return ""

        deadline = time.monotonic() + self.options.ocr_timeout_seconds
        with TemporaryDirectory(prefix="papagui-ocr-") as temp_dir:
            output_prefix = Path(temp_dir) / "page"
            render = subprocess.run(
                [
                    pdftoppm,
                    "-jpeg",
                    "-jpegopt",
                    "quality=85",
                    "-r",
                    "120",
                    "-f",
                    "1",
                    "-l",
                    str(self.options.ocr_max_pages),
                    str(filepath),
                    str(output_prefix),
                ],
                capture_output=True,
                text=True,
                timeout=max(5, min(15, self.options.ocr_timeout_seconds)),
            )
            if render.returncode != 0:
                return ""

            image_paths = sorted(Path(temp_dir).glob("page-*.jpg"))
            language = self._get_ocr_language(tesseract)

            def recognize(image_path: Path) -> tuple[Path, str]:
                remaining = max(1, deadline - time.monotonic())
                result = subprocess.run(
                    [
                        str(tesseract),
                        str(image_path),
                        "stdout",
                        "-l",
                        language,
                    ],
                    capture_output=True,
                    text=True,
                    errors="replace",
                    timeout=max(1, min(15, remaining)),
                )
                return image_path, result.stdout if result.returncode == 0 else ""

            recognized = {}
            with ThreadPoolExecutor(max_workers=min(4, max(1, len(image_paths)))) as executor:
                futures = [executor.submit(recognize, image_path) for image_path in image_paths]
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise TimeoutError(
                        f"OCR-Zeitlimit von {self.options.ocr_timeout_seconds} Sekunden erreicht"
                    )
                try:
                    for future in as_completed(futures, timeout=remaining):
                        image_path, text = future.result()
                        if text.strip():
                            recognized[image_path] = text
                except TimeoutError as exc:
                    raise TimeoutError(
                        f"OCR-Zeitlimit von {self.options.ocr_timeout_seconds} Sekunden erreicht"
                    ) from exc
            return self._limit_text(
                "\n".join(recognized[path] for path in image_paths if path in recognized)
            )

    def _get_ocr_language(self, tesseract: Path) -> str:
        if self._ocr_language is not None:
            return self._ocr_language
        try:
            result = subprocess.run(
                [str(tesseract), "--list-langs"],
                capture_output=True,
                text=True,
                timeout=10,
            )
            languages = set(result.stdout.splitlines()[1:])
        except Exception:
            languages = set()
        if "deu" in languages and "eng" in languages:
            self._ocr_language = "deu+eng"
        elif "deu" in languages:
            self._ocr_language = "deu"
        else:
            self._ocr_language = "eng"
        return self._ocr_language

    def _find_tesseract(self) -> Optional[Path]:
        found = shutil.which("tesseract")
        if found:
            return Path(found)
        if sys.platform == "win32":
            for environment_name in ("PROGRAMFILES", "LOCALAPPDATA"):
                base = os.environ.get(environment_name)
                if base:
                    candidate = Path(base) / "Tesseract-OCR" / "tesseract.exe"
                    if candidate.exists():
                        return candidate
        return None

    def _find_libreoffice(self) -> Optional[Path]:
        for name in ("libreoffice", "soffice"):
            found = shutil.which(name)
            if found:
                return Path(found)

        candidates = []
        if sys.platform == "darwin":
            candidates.append(Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"))
        elif sys.platform == "win32":
            for environment_name in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
                base = os.environ.get(environment_name)
                if base:
                    candidates.append(Path(base) / "LibreOffice" / "program" / "soffice.exe")
        return next((path for path in candidates if path.exists()), None)
    
    def search_customers(self, query: str, limit: Optional[int] = None) -> List[Dict]:
        cursor = self.conn.cursor()
        
        q = f"%{query}%"
        sql = f"""
            SELECT
                {self.SEARCH_LABEL_SQL} AS customer_name,
                domain_folder AS service_type,
                time_bucket AS year,
                COUNT(*) AS file_count
            FROM files
            WHERE (
                {self.SEARCH_LABEL_SQL} LIKE ?
                OR relative_dir LIKE ?
                OR domain_folder LIKE ?
            )
            GROUP BY customer_name, service_type, year
            ORDER BY customer_name
        """
        params = [q, q, q]
        if limit is not None:
            sql += " LIMIT ?"
            params.append(limit)
        cursor.execute(sql, params)
        
        results = []
        for row in cursor.fetchall():
            results.append({
                'customer_name': row['customer_name'],
                'service_type': row['service_type'],
                'year': row['year'],
                'file_count': row['file_count']
            })
        
        return results

    def search_folders(self, query: str, limit: Optional[int] = None) -> List[Dict]:
        return self.search_folders_page(query, SearchFilters(), 1, limit or 200).items

    def _metadata_filter_clause(
        self, filters: SearchFilters, alias: str = "files"
    ) -> tuple[str, list[str]]:
        clauses = []
        params = []
        if filters.domain_folder:
            clauses.append(f"{alias}.domain_folder = ? COLLATE NOCASE")
            params.append(filters.domain_folder)
        if filters.year:
            clauses.append(f"{alias}.time_bucket = ? COLLATE NOCASE")
            params.append(filters.year)
        if filters.file_type:
            clauses.append(f"{alias}.file_type = ? COLLATE NOCASE")
            params.append(filters.file_type.lstrip("."))
        return (" AND " + " AND ".join(clauses) if clauses else "", params)

    def get_search_facets(self) -> Dict[str, List[str]]:
        values = {}
        for key, column in (
            ("domains", "domain_folder"),
            ("years", "time_bucket"),
            ("file_types", "file_type"),
        ):
            rows = self.conn.execute(
                f"""
                SELECT DISTINCT {column} FROM files
                WHERE COALESCE({column}, '') <> ''
                ORDER BY {column} COLLATE NOCASE
                """
            )
            values[key] = [str(row[0]) for row in rows]
        values["years"] = sorted(
            values["years"], key=lambda value: (not value.isdigit(), value), reverse=True
        )
        return values

    def search_folders_page(
        self,
        query: str,
        filters: SearchFilters,
        page: int = 1,
        page_size: int = 25,
    ) -> SearchPage:
        cursor = self.conn.cursor()
        filter_sql, filter_params = self._metadata_filter_clause(filters, "f")
        exists_sql = f"""
            EXISTS (
                SELECT 1 FROM files f
                WHERE (f.folder_path = folders.path OR f.path LIKE folders.path || ? || '%')
                {filter_sql}
            )
        """
        where_sql = f"(folders.name LIKE ? OR folders.relative_path LIKE ?) AND {exists_sql}"
        common_params = [f"%{query}%", f"%{query}%", os.sep, *filter_params]
        total = int(cursor.execute(
            f"SELECT COUNT(*) FROM folders WHERE {where_sql}", common_params
        ).fetchone()[0])
        offset = max(0, page - 1) * page_size
        sql = f"""
            SELECT path, name, relative_path,
                (SELECT COUNT(*) FROM files
                 WHERE files.folder_path = folders.path
                    OR files.path LIKE folders.path || ? || '%') AS file_count
            FROM folders
            WHERE {where_sql}
            ORDER BY
                CASE
                    WHEN name = ? COLLATE NOCASE THEN 0
                    WHEN name LIKE ? THEN 1
                    WHEN name LIKE ? THEN 2
                    ELSE 3
                END,
                name COLLATE NOCASE, relative_path COLLATE NOCASE
            LIMIT ? OFFSET ?
        """
        params = [
            os.sep,
            *common_params,
            query,
            f"{query}%",
            f"%{query}%",
            page_size,
            offset,
        ]
        cursor.execute(sql, params)
        items = [
            {
                "folder_path": row["path"],
                "folder_name": row["name"],
                "relative_path": row["relative_path"],
                "file_count": row["file_count"],
            }
            for row in cursor.fetchall()
        ]
        return SearchPage(items, total, page, page_size)

    def get_folder_search_entry(
        self, folder_path: str, filters: SearchFilters
    ) -> Optional[Dict]:
        filter_sql, filter_params = self._metadata_filter_clause(filters, "f")
        row = self.conn.execute(
            f"""
            SELECT folders.path, folders.name, folders.relative_path,
                (SELECT COUNT(*) FROM files
                 WHERE files.folder_path = folders.path
                    OR files.path LIKE folders.path || ? || '%') AS file_count
            FROM folders
            WHERE folders.path = ?
              AND EXISTS (
                SELECT 1 FROM files f
                WHERE (f.folder_path = folders.path OR f.path LIKE folders.path || ? || '%')
                {filter_sql}
              )
            """,
            (os.sep, str(Path(folder_path)), os.sep, *filter_params),
        ).fetchone()
        if row is None:
            return None
        return {
            "folder_path": row["path"],
            "folder_name": row["name"],
            "relative_path": row["relative_path"],
            "file_count": row["file_count"],
        }

    def get_folder_details(self, folder_path: str) -> Dict:
        cursor = self.conn.cursor()
        path = str(Path(folder_path))
        pattern = f"{path}{os.sep}%"
        cursor.execute(
            """
            SELECT COUNT(*) AS file_count, MAX(modified_date) AS last_modified,
                   COALESCE(SUM(file_size), 0) AS total_size
            FROM files WHERE folder_path = ? OR path LIKE ?
            """,
            (path, pattern),
        )
        info = dict(cursor.fetchone())
        cursor.execute(
            """
            SELECT path, filename, file_type, file_size, modified_date,
                   domain_folder, time_bucket, relative_dir
            FROM files WHERE folder_path = ? OR path LIKE ?
            ORDER BY relative_dir, filename COLLATE NOCASE
            """,
            (path, pattern),
        )
        info["files"] = [dict(row) for row in cursor.fetchall()]
        info["folder_name"] = Path(path).name
        info["folder_path"] = path
        info["service_types"] = sorted({
            row["domain_folder"] for row in info["files"] if row["domain_folder"]
        })
        return info
    
    def get_customer_details(self, customer_name: str) -> Dict:
        cursor = self.conn.cursor()

        cursor.execute(
            f"""
            SELECT 
                COUNT(*) as file_count,
                MAX(modified_date) as last_modified,
                SUM(file_size) as total_size
            FROM files
            WHERE {self.SEARCH_LABEL_SQL} = ?
            """,
            (customer_name,)
        )
        
        info = dict(cursor.fetchone())
        
        cursor.execute(
            f"""
            SELECT DISTINCT domain_folder FROM files
            WHERE {self.SEARCH_LABEL_SQL} = ?
            ORDER BY domain_folder
            """,
            (customer_name,)
        )
        
        info['service_types'] = [row['domain_folder'] for row in cursor.fetchall() if row['domain_folder']]
        
        cursor.execute(
            f"""
            SELECT 
                path, filename, file_type, file_size, modified_date,
                domain_folder, time_bucket, project_name, subfolder, relative_dir,
                service_type, year
            FROM files
            WHERE {self.SEARCH_LABEL_SQL} = ?
            ORDER BY time_bucket DESC, domain_folder, subfolder, filename
            """,
            (customer_name,)
        )
        
        info['files'] = [dict(row) for row in cursor.fetchall()]
        info['customer_name'] = customer_name
        
        return info
    
    def search_files(
        self,
        query: str,
        customer_name: Optional[str] = None,
        limit: Optional[int] = None,
    ) -> List[Dict]:
        return self.search_files_page(
            query, SearchFilters(), 1, limit or 200, customer_name=customer_name
        ).items

    def search_files_page(
        self,
        query: str,
        filters: SearchFilters,
        page: int = 1,
        page_size: int = 25,
        customer_name: Optional[str] = None,
    ) -> SearchPage:
        cursor = self.conn.cursor()
        filter_sql, filter_params = self._metadata_filter_clause(filters)
        where = "(filename LIKE ? OR relative_dir LIKE ?)" + filter_sql
        params: list = [f"%{query}%", f"%{query}%", *filter_params]
        if customer_name:
            where += f" AND {self.SEARCH_LABEL_SQL} = ?"
            params.append(customer_name)
        total = int(cursor.execute(f"SELECT COUNT(*) FROM files WHERE {where}", params).fetchone()[0])
        offset = max(0, page - 1) * page_size
        sql = f"""
            SELECT * FROM files WHERE {where}
            ORDER BY
                CASE
                    WHEN filename = ? COLLATE NOCASE THEN 0
                    WHEN filename LIKE ? THEN 1
                    WHEN filename LIKE ? THEN 2
                    ELSE 3
                END,
                modified_date DESC, filename COLLATE NOCASE
            LIMIT ? OFFSET ?
        """
        cursor.execute(
            sql,
            [*params, query, f"{query}%", f"%{query}%", page_size, offset],
        )
        return SearchPage([dict(row) for row in cursor.fetchall()], total, page, page_size)

    def _iter_text_search_targets(self, customer_name: Optional[str]) -> List[List[str]]:
        cursor = self.conn.cursor()

        if customer_name:
            cursor.execute(
                f"SELECT DISTINCT path FROM files WHERE {self.SEARCH_LABEL_SQL} = ?",
                (customer_name,)
            )
            paths = [row[0] for row in cursor.fetchall()]
            chunk_size = 200
            return [paths[i:i + chunk_size] for i in range(0, len(paths), chunk_size)]

        cursor.execute("SELECT DISTINCT root_path FROM indexed_roots")
        roots = [row[0] for row in cursor.fetchall()]
        return [[root] for root in roots]

    def _parse_rg_output(self, stdout: str) -> List[Dict]:
        parsed = []
        for line in stdout.splitlines():
            try:
                event = json.loads(line)
            except json.JSONDecodeError:
                continue
            if event.get("type") != "match":
                continue
            data = event.get("data", {})
            path_data = data.get("path", {})
            lines_data = data.get("lines", {})
            file_path = path_data.get("text")
            if not file_path:
                continue
            parsed.append({
                "path": file_path,
                "line": data.get("line_number"),
                "excerpt": (lines_data.get("text") or "").strip(),
                "source": "text",
            })
        return parsed

    def _build_fts_query(self, query: str) -> str:
        tokens = re.findall(r"\w+", query, flags=re.UNICODE)
        return " AND ".join(f'"{token.replace(chr(34), chr(34) * 2)}"*' for token in tokens)

    def _search_extracted_content(
        self,
        query: str,
        customer_name: Optional[str],
        limit: int = 200,
        filters: Optional[SearchFilters] = None,
    ) -> List[Dict]:
        fts_query = self._build_fts_query(query)
        if not fts_query:
            return []

        sql = f"""
            SELECT
                file_content_fts.path AS path,
                snippet(file_content_fts, 1, '', '', ' … ', 18) AS excerpt
            FROM file_content_fts
            JOIN files ON files.path = file_content_fts.path
            WHERE file_content_fts MATCH ?
        """
        params = [fts_query]
        filter_sql, filter_params = self._metadata_filter_clause(filters or SearchFilters())
        sql += filter_sql
        params.extend(filter_params)
        if customer_name:
            sql += f" AND {self.SEARCH_LABEL_SQL} = ?"
            params.append(customer_name)
        sql += " ORDER BY rank LIMIT ?"
        params.append(limit)

        cursor = self.conn.cursor()
        cursor.execute(sql, params)
        return [
            {
                "path": row["path"],
                "line": None,
                "excerpt": (row["excerpt"] or "").replace("\n", " ").strip(),
                "source": "document",
            }
            for row in cursor.fetchall()
        ]
    
    def _run_ripgrep(
        self,
        command: List[str],
        should_cancel: Optional[Callable[[], bool]],
    ) -> str:
        process = subprocess.Popen(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            text=True,
        )
        deadline = time.monotonic() + 25
        while True:
            try:
                stdout, _ = process.communicate(timeout=0.1)
                return stdout
            except subprocess.TimeoutExpired:
                if should_cancel is not None and should_cancel():
                    process.terminate()
                    try:
                        process.communicate(timeout=1)
                    except subprocess.TimeoutExpired:
                        process.kill()
                        process.communicate()
                    return ""
                if time.monotonic() >= deadline:
                    process.kill()
                    process.communicate()
                    return ""

    def search_in_text(
        self,
        query: str,
        customer_name: Optional[str] = None,
        should_cancel: Optional[Callable[[], bool]] = None,
        limit: int = 200,
        filters: Optional[SearchFilters] = None,
    ) -> List[Dict]:
        try:
            results = self._search_extracted_content(
                query, customer_name, limit=limit, filters=filters
            )
            indexed_result_paths = {result["path"] for result in results}
            if should_cancel is not None and should_cancel():
                return results
            if filters is not None and filters.active:
                return results

            rg_path = shutil.which("rg")
            if rg_path is None:
                return results

            target_groups = self._iter_text_search_targets(customer_name)
            if not target_groups:
                return results

            for targets in target_groups:
                if should_cancel is not None and should_cancel():
                    break
                command = [
                    rg_path,
                    '--json',
                    '--line-number',
                    '--no-messages',
                    '--smart-case',
                    '--fixed-strings',
                    '--glob', '*.txt',
                    '--glob', '*.csv',
                    '--glob', '*.md',
                    '--glob', '*.log',
                    '--glob', '*.json',
                    '--glob', '*.xml',
                    '--glob', '*.yaml',
                    '--glob', '*.yml',
                    '--glob', '*.ini',
                    query,
                    *targets,
                ]
                stdout = self._run_ripgrep(command, should_cancel)
                if stdout:
                    results.extend(
                        result
                        for result in self._parse_rg_output(stdout)
                        if result["path"] not in indexed_result_paths
                    )

            unique_results = []
            seen = set()
            for result in results:
                key = (result["path"], result.get("line"))
                if key in seen:
                    continue
                seen.add(key)
                unique_results.append(result)
                if len(unique_results) >= limit:
                    break
            return unique_results
        except Exception:
            logger.exception("Fehler bei Volltextsuche: query=%r", query)
            return []

    def search_text_page(
        self,
        query: str,
        filters: SearchFilters,
        page: int = 1,
        page_size: int = 25,
        maximum: int = 200,
        should_cancel: Optional[Callable[[], bool]] = None,
    ) -> SearchPage:
        all_results = self.search_in_text(
            query,
            should_cancel=should_cancel,
            limit=maximum,
            filters=filters,
        )
        offset = max(0, page - 1) * page_size
        return SearchPage(
            all_results[offset:offset + page_size],
            len(all_results),
            page,
            page_size,
        )
    
    def close(self):
        """Schließt die Datenbankverbindung"""
        if self.conn:
            self.conn.close()


if __name__ == "__main__":
    from app.core.config import DB_FILE, get_default_index_source
    
    manager = IndexManager(DB_FILE)
    index_root = get_default_index_source()
    manager.index_directory(index_root)
    
    results = manager.search_customers("Müller")
    logger.info("Gefundene Kunden: %s", results)
    
    manager.close()
