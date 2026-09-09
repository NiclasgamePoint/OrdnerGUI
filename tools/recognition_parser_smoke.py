"""Synthetic runtime parser/OCR smoke; creates and destroys all input locally.

Example: python tools/recognition_parser_smoke.py
No existing documents, databases, network services or customer paths are opened.
"""

from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
import openpyxl
from PIL import Image, ImageDraw, ImageFont

from papagui_server.adapters.catalog_extraction import DocumentTextExtractor
from papagui_server.domain.models import ServerSettings


def _scan():
    image = Image.new("RGB", (1600, 2000), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=52)
    lines = (
        "SCANNED SYNTHETIC CONTACT",
        "Example Company GmbH",
        "Phone: +49 30 12345678",
        "Email: synthetic@example.org",
        "Example Street 12",
        "01234 Example City",
        "This document is generated for parser testing.",
        "It contains no real customer information.",
    )
    for index, line in enumerate(lines):
        draw.text((100, 120 + index * 100), line, fill="black", font=font)
    return image


def _mixed_pdf(path: Path, scan) -> None:
    """Create two real PDF pages without requiring a test-only PDF library."""
    encoded = BytesIO()
    scan.save(encoded, format="JPEG", quality=95)
    jpeg = encoded.getvalue()
    native_lines = [
        "NATIVE SYNTHETIC CONTACT",
        "Example Company GmbH",
        "Phone +49 30 12345678",
        "Generated exclusively for synthetic parser verification.",
    ]
    native = (
        b"BT /F1 14 Tf 40 740 Td "
        + b" 0 -25 Td ".join(b"(" + line.encode("ascii") + b") Tj" for line in native_lines)
        + b" ET"
    )
    scanned = b"q 540 0 0 675 35 65 cm /Im1 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 6 0 R >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /XObject << /Im1 7 0 R >> >> /Contents 8 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(native)).encode() + b" >>\nstream\n" + native + b"\nendstream",
        b"<< /Type /XObject /Subtype /Image /Width 1600 /Height 2000 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length "
        + str(len(jpeg)).encode()
        + b" >>\nstream\n"
        + jpeg
        + b"\nendstream",
        b"<< /Length " + str(len(scanned)).encode() + b" >>\nstream\n" + scanned + b"\nendstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(str(number).encode() + b" 0 obj\n" + value + b"\nendobj\n")
    xref = len(result)
    result.extend(b"xref\n0 9\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(
        b"trailer\n<< /Size 9 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF\n"
    )
    path.write_bytes(result)


def run() -> dict:
    settings = ServerSettings(
        ocr_max_pages=3,
        ocr_extended_max_pages=3,
        ocr_extension_threshold=100,
        ocr_timeout_seconds=30,
        extraction_timeout_seconds=90,
        max_extracted_characters=10000,
    )
    extractor = DocumentTextExtractor()
    results = {}
    with TemporaryDirectory(prefix="papagui-synthetic-parser-smoke-") as temporary:
        root = Path(temporary)
        word = Document()
        word.sections[0].header.paragraphs[0].text = "SYNTHETIC HEADER"
        paragraph = word.add_paragraph()
        paragraph.add_run("Example ")
        paragraph.add_run("Company GmbH")
        word.sections[0].footer.paragraphs[0].text = "SYNTHETIC FOOTER"
        word.save(root / "synthetic.docx")
        result = extractor.extract_document(root / "synthetic.docx", settings)
        assert result.status == "ok" and "Example Company GmbH" in result.text, "docx_text"
        assert "SYNTHETIC HEADER" in result.text and "SYNTHETIC FOOTER" in result.text, (
            "docx_sections"
        )
        results["docx"] = result.status
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Synthetic"
        sheet["A1"] = 1234
        sheet["A1"].number_format = "00000"
        workbook.save(root / "synthetic.xlsx")
        workbook.close()
        result = extractor.extract_document(root / "synthetic.xlsx", settings)
        assert (
            result.status == "ok"
            and result.blocks[0].text == "01234"
            and result.blocks[0].cell == "A1"
        ), "xlsx_values"
        results["xlsx"] = result.status
        scan = _scan()
        scan.save(root / "scan.png")
        result = extractor.extract_document(root / "scan.png", settings)
        assert result.status == "ok" and "SYNTHETIC" in result.text, "image_ocr"
        assert any(block.method == "tesseract_tsv" and block.bbox for block in result.blocks), (
            "image_layout"
        )
        results["image_ocr"] = result.status
        scan.rotate(90, expand=True).save(root / "scan_rotated.png")
        result = extractor.extract_document(root / "scan_rotated.png", settings)
        assert result.status == "ok" and "SYNTHETIC" in result.text, "rotated_image_ocr"
        results["rotated_image_ocr"] = result.status
        _mixed_pdf(root / "synthetic.pdf", scan)
        result = extractor.extract_document(root / "synthetic.pdf", settings)
        assert result.status == "ok" and result.pages_total == result.pages_processed == 2, (
            "mixed_pdf_pages"
        )
        assert "NATIVE SYNTHETIC" in result.text and "SCANNED SYNTHETIC" in result.text, (
            "mixed_pdf_text"
        )
        assert {block.method for block in result.blocks} == {"pdfplumber", "tesseract_tsv"}, (
            "mixed_pdf_layout"
        )
        results["mixed_pdf"] = result.status
    return {"status": "passed", "synthetic": True, "checks": results}


if __name__ == "__main__":
    try:
        print(json.dumps(run(), sort_keys=True))
    except AssertionError as error:
        print(json.dumps({"status": "failed", "synthetic": True, "check": str(error)}))
        raise SystemExit(1) from None
