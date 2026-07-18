from __future__ import annotations

from pathlib import Path

from docx import Document
from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import QDesktopServices, QPixmap, QWheelEvent
from PySide6.QtWidgets import (
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QScrollArea,
    QStackedWidget,
    QVBoxLayout,
    QWidget,
)

from app.gui.viewers import PdfViewerWidget, SpreadsheetViewerWidget, TextViewerWidget
from app.gui.widgets.buttons import AppButton, BusyIndicator
from app.gui.workers.file_conversion_worker import FileConversionWorker
from app.services.document_converter import DocumentConverter


class ImageViewerWidget(QScrollArea):
    """Display images fitted to viewport with optional Ctrl+wheel zoom."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("ImageViewerScroll")
        self.setWidgetResizable(False)
        self.setAlignment(Qt.AlignCenter)
        self._image_label = QLabel()
        self._image_label.setAlignment(Qt.AlignCenter)
        self.setWidget(self._image_label)
        self._source_pixmap: QPixmap | None = None
        self._zoom_multiplier = 1.0

    def set_image(self, pixmap: QPixmap):
        self._source_pixmap = pixmap
        self._zoom_multiplier = 1.0
        self._update_display_pixmap()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self._source_pixmap is not None:
            self._update_display_pixmap()

    def wheelEvent(self, event: QWheelEvent):
        if event.modifiers() & Qt.ControlModifier and self._source_pixmap is not None:
            delta = event.angleDelta().y()
            if delta != 0:
                factor = 1.15 if delta > 0 else 1 / 1.15
                self._zoom_multiplier = max(1.0, min(8.0, self._zoom_multiplier * factor))
                self._update_display_pixmap()
            event.accept()
            return
        super().wheelEvent(event)

    def _update_display_pixmap(self):
        if self._source_pixmap is None:
            return
        source_width = self._source_pixmap.width()
        source_height = self._source_pixmap.height()
        if source_width <= 0 or source_height <= 0:
            return

        viewport_size = self.viewport().size()
        viewport_width = max(1, viewport_size.width())
        viewport_height = max(1, viewport_size.height())
        fit_scale = min(viewport_width / source_width, viewport_height / source_height, 1.0)
        scale = fit_scale * self._zoom_multiplier

        scaled = self._source_pixmap.scaled(
            max(1, int(source_width * scale)),
            max(1, int(source_height * scale)),
            Qt.KeepAspectRatio,
            Qt.SmoothTransformation,
        )
        self._image_label.setPixmap(scaled)
        self._image_label.resize(scaled.size())


class FileViewer(QWidget):
    """Route supported formats to focused viewer widgets."""

    TEXT_TYPES = {"txt", "csv", "log", "md", "json", "xml", "yaml", "yml", "ini"}
    IMAGE_TYPES = {"jpg", "jpeg", "png", "gif", "bmp", "webp"}

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("FileViewer")
        self.current_file: Path | None = None
        self._load_generation = 0
        self._conversion_workers: set[FileConversionWorker] = set()
        self._active_converter: DocumentConverter | None = None
        self._word_preview_error = ""

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        header = QHBoxLayout()
        self.file_info_label = QLabel("Keine Datei geladen")
        self.file_info_label.setObjectName("ViewerFileInfo")
        header.addWidget(self.file_info_label, 1)
        self.open_folder_button = AppButton("Ordner", AppButton.SECONDARY)
        self.open_external_button = AppButton("Extern öffnen", AppButton.SECONDARY)
        self.open_folder_button.clicked.connect(self.open_containing_folder)
        self.open_external_button.clicked.connect(self.open_externally)
        self.open_folder_button.setEnabled(False)
        self.open_external_button.setEnabled(False)
        header.addWidget(self.open_folder_button)
        header.addWidget(self.open_external_button)
        layout.addLayout(header)

        self.stack = QStackedWidget()
        self.stack.setObjectName("ViewerStack")
        self.empty_label = QLabel("Keine Datei geladen")
        self.empty_label.setAlignment(Qt.AlignCenter)
        self.loading_widget = QWidget()
        self.loading_widget.setObjectName("ViewerLoading")
        loading_layout = QVBoxLayout(self.loading_widget)
        loading_layout.setContentsMargins(24, 24, 24, 24)
        loading_layout.addStretch(1)
        self.loading_indicator = BusyIndicator()
        self.loading_indicator.setAlignment(Qt.AlignCenter)
        self.loading_label = QLabel("Datei wird vorbereitet …")
        self.loading_label.setObjectName("ViewerLoadingText")
        self.loading_label.setAlignment(Qt.AlignCenter)
        self.loading_label.setWordWrap(True)
        loading_layout.addWidget(self.loading_indicator, 0, Qt.AlignCenter)
        loading_layout.addWidget(self.loading_label)
        loading_layout.addStretch(1)
        self.pdf_viewer = PdfViewerWidget()
        self.spreadsheet_viewer = SpreadsheetViewerWidget()
        self.text_viewer = TextViewerWidget()
        self.image_viewer = ImageViewerWidget()
        for widget in (
            self.empty_label,
            self.loading_widget,
            self.pdf_viewer,
            self.spreadsheet_viewer,
            self.text_viewer,
            self.image_viewer,
        ):
            self.stack.addWidget(widget)
        layout.addWidget(self.stack, 1)
        self.conversion_meta_label = QLabel("")
        self.conversion_meta_label.setObjectName("ViewerMeta")
        self.conversion_meta_label.setWordWrap(True)
        self.conversion_meta_label.setVisible(False)
        layout.addWidget(self.conversion_meta_label)

    def open_file(self, filepath: Path):
        filepath = filepath.resolve()
        if not filepath.exists() or not filepath.is_file():
            QMessageBox.warning(self, "Fehler", f"Datei nicht gefunden: {filepath}")
            return
        self._load_generation += 1
        self._cancel_conversions()
        self._cleanup_active_converter()
        self._word_preview_error = ""
        self._stop_loading()
        self.pdf_viewer.close_document()
        self.current_file = filepath
        self.file_info_label.setText(
            f"{filepath.name} · {filepath.stat().st_size / 1024:.1f} KB"
        )
        self.open_folder_button.setEnabled(True)
        self.open_external_button.setEnabled(True)
        self._set_conversion_meta(converted=False, tool="Direkt")
        suffix = filepath.suffix.lower().lstrip(".")
        try:
            if suffix == "pdf":
                self._show_pdf(filepath)
            elif suffix in {"xlsx", "xls"}:
                self._show_spreadsheet(filepath)
            elif suffix in self.TEXT_TYPES:
                self._show_text(filepath.read_text(encoding="utf-8", errors="replace"))
            elif suffix == "docx":
                self._show_word_preview(filepath)
            elif suffix == "doc":
                self._show_word_preview(filepath)
            elif suffix in self.IMAGE_TYPES:
                self._show_image(filepath)
            else:
                self._show_text(
                    f"Dateiformat: {filepath.suffix or 'ohne Endung'}\n\n"
                    "Für dieses Format steht keine interne Vorschau zur Verfügung. "
                    "Die Datei kann über ‚Extern öffnen‘ angezeigt werden."
                )
        except Exception as exc:
            self._show_text(f"Datei konnte nicht angezeigt werden:\n\n{exc}")

    def _show_pdf(self, path: Path):
        self.pdf_viewer.load(path)
        self.stack.setCurrentWidget(self.pdf_viewer)
        self._set_conversion_meta(converted=False, tool="Direkt")

    def _show_spreadsheet(self, path: Path):
        try:
            self.spreadsheet_viewer.load(path)
        except Exception:
            if path.suffix.lower() != ".xls":
                raise
            self._start_conversion(
                FileConversionWorker.XLS_TO_XLSX,
                path,
                "Excel-Datei wird konvertiert …",
            )
            return
        self.stack.setCurrentWidget(self.spreadsheet_viewer)
        self._set_conversion_meta(converted=False, tool="Direkt")

    def _show_docx_text(self, path: Path, prefix: str = ""):
        document = Document(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(lines)
        if prefix:
            text = f"{prefix}\n\n{text}"
        self._show_text(text)
        self._set_conversion_meta(converted=False, tool="Direkt")

    def _show_word_preview(self, path: Path):
        self._start_conversion(
            FileConversionWorker.WORD_TO_PDF,
            path,
            "Word-Datei wird als formatierte Vorschau vorbereitet …",
        )

    def _show_legacy_doc(self, path: Path):
        self._start_conversion(
            FileConversionWorker.EXTRACT_DOC,
            path,
            "Word-Datei wird als Textvorschau verarbeitet …",
        )

    def _start_conversion(self, operation: str, path: Path, message: str):
        self._start_loading(message)
        worker = FileConversionWorker(
            self._load_generation,
            operation,
            path,
            parent=self,
        )
        worker.completed.connect(self._on_conversion_completed)
        worker.finished.connect(
            lambda worker=worker: self._release_conversion_worker(worker)
        )
        self._conversion_workers.add(worker)
        worker.start()

    def _on_conversion_completed(
        self,
        generation: int,
        operation: str,
        result,
        metadata,
        error: str,
    ):
        worker = self.sender()
        try:
            if generation != self._load_generation:
                return
            self._stop_loading()
            if error:
                if error != "abgebrochen":
                    self._handle_conversion_error(operation, error)
                return

            if operation == FileConversionWorker.XLS_TO_XLSX:
                self.spreadsheet_viewer.load(Path(result))
                if isinstance(worker, FileConversionWorker):
                    self._active_converter = worker.take_converter()
                self.stack.setCurrentWidget(self.spreadsheet_viewer)
            elif operation == FileConversionWorker.WORD_TO_PDF:
                self.pdf_viewer.load(Path(result))
                if isinstance(worker, FileConversionWorker):
                    self._active_converter = worker.take_converter()
                self.stack.setCurrentWidget(self.pdf_viewer)
            elif operation == FileConversionWorker.EXTRACT_DOC:
                text = str(result or "")
                if not text.strip():
                    raise ValueError(
                        "Inhalt der .doc-Datei konnte nicht extrahiert werden"
                    )
                if self._word_preview_error:
                    text = (
                        "Formatierte Word-Vorschau konnte nicht erstellt werden:\n"
                        f"{self._word_preview_error}\n\nTextinhalt:\n\n{text}"
                    )
                    self._word_preview_error = ""
                self._show_text(text)
            self._set_conversion_meta(
                converted=bool(metadata.get("converted", False)),
                tool=str(metadata.get("tool", "Unbekannt")),
            )
        except Exception as exc:
            if generation == self._load_generation:
                self._stop_loading()
                self._set_conversion_meta(converted=False, tool="")
                self._show_text(
                    "Datei konnte nicht angezeigt werden:\n\n"
                    f"{exc}"
                )
        finally:
            if isinstance(worker, FileConversionWorker):
                worker.cleanup()

    def _handle_conversion_error(self, operation: str, error: str):
        if operation == FileConversionWorker.WORD_TO_PDF and self.current_file is not None:
            suffix = self.current_file.suffix.lower()
            self._word_preview_error = error
            if suffix == ".docx":
                try:
                    self._show_docx_text(
                        self.current_file,
                        prefix=(
                            "Formatierte Word-Vorschau konnte nicht erstellt werden:\n"
                            f"{error}\n\nTextinhalt:"
                        ),
                    )
                    return
                except Exception as fallback_error:
                    self._set_conversion_meta(converted=False, tool="")
                    self._show_text(
                        "Datei konnte nicht angezeigt werden:\n\n"
                        f"{fallback_error}"
                    )
                    return
            if suffix == ".doc":
                self._show_legacy_doc(self.current_file)
                return
        self._set_conversion_meta(converted=False, tool="")
        self._show_text(
            "Datei konnte nicht konvertiert werden:\n\n"
            f"{error}"
        )

    def _release_conversion_worker(self, worker: FileConversionWorker):
        self._conversion_workers.discard(worker)
        worker.deleteLater()

    def _start_loading(self, message: str):
        self.loading_label.setText(message)
        self.loading_indicator.start()
        self.stack.setCurrentWidget(self.loading_widget)
        self.conversion_meta_label.setText(
            "Die Konvertierung läuft unabhängig von der Oberfläche."
        )
        self.conversion_meta_label.setVisible(True)

    def _stop_loading(self):
        self.loading_indicator.stop()

    def _cancel_conversions(self):
        for worker in tuple(self._conversion_workers):
            if worker.isRunning():
                worker.requestInterruption()

    def _cleanup_active_converter(self):
        if self._active_converter is not None:
            self._active_converter.cleanup()
            self._active_converter = None

    @property
    def is_converting(self) -> bool:
        return any(worker.isRunning() for worker in self._conversion_workers)

    def _show_text(self, text: str):
        self.text_viewer.set_text(text[:2_000_000])
        self.stack.setCurrentWidget(self.text_viewer)

    def _show_image(self, path: Path):
        pixmap = QPixmap(str(path))
        if pixmap.isNull():
            raise ValueError("Bild konnte nicht geladen werden")
        self.image_viewer.set_image(pixmap)
        self.stack.setCurrentWidget(self.image_viewer)
        self._set_conversion_meta(converted=False, tool="Direkt")

    def _set_conversion_meta(self, *, converted: bool, tool: str):
        if converted:
            self.conversion_meta_label.setText(f"Konvertiert mit: {tool}")
            self.conversion_meta_label.setVisible(True)
            return
        self.conversion_meta_label.clear()
        self.conversion_meta_label.setVisible(False)

    def open_externally(self):
        if self.current_file is not None:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.current_file)))

    def open_containing_folder(self):
        if self.current_file is not None:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.current_file.parent)))

    def closeEvent(self, event):
        self.shutdown()
        super().closeEvent(event)

    def shutdown(self):
        self._load_generation += 1
        self._stop_loading()
        self._cancel_conversions()
        for worker in tuple(self._conversion_workers):
            worker.wait()
            worker.cleanup()
        self._conversion_workers.clear()
        self._cleanup_active_converter()
        self.pdf_viewer.close_document()
