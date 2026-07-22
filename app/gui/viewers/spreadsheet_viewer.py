from __future__ import annotations

from pathlib import Path

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter
from PySide6.QtWidgets import QComboBox, QHBoxLayout, QLabel, QLineEdit, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget

from app.gui.widgets.buttons import AppButton


class SpreadsheetViewerWidget(QWidget):
    MAX_ROWS = 500
    MAX_COLUMNS = 50

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("ViewerContent")
        self.path: Path | None = None
        self.sheet_names: list[str] = []
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        toolbar = QHBoxLayout()
        toolbar.addWidget(QLabel("Tabellenblatt:"))
        self.sheet_combo = QComboBox()
        self.sheet_combo.currentIndexChanged.connect(self._load_current_sheet)
        toolbar.addWidget(self.sheet_combo)
        toolbar.addStretch()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("In Tabelle suchen …")
        self.search_input.setMaximumWidth(210)
        self.search_button = AppButton("Weiter", AppButton.SECONDARY)
        self.search_input.returnPressed.connect(self.find_next)
        self.search_button.clicked.connect(self.find_next)
        toolbar.addWidget(self.search_input)
        toolbar.addWidget(self.search_button)
        layout.addLayout(toolbar)
        self.table = QTableWidget()
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        layout.addWidget(self.table, 1)
        self.info_label = QLabel()
        self.info_label.setObjectName("ViewerMeta")
        layout.addWidget(self.info_label)

    def load(self, path: Path):
        self.path = path
        if path.suffix.lower() == ".xlsx":
            workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            try:
                self.sheet_names = list(workbook.sheetnames)
            finally:
                workbook.close()
        else:
            workbook = xlrd.open_workbook(str(path), on_demand=True)
            try:
                self.sheet_names = workbook.sheet_names()
            finally:
                workbook.release_resources()
        self.sheet_combo.blockSignals(True)
        self.sheet_combo.clear()
        self.sheet_combo.addItems(self.sheet_names)
        self.sheet_combo.blockSignals(False)
        self._load_current_sheet()

    def _load_current_sheet(self):
        if self.path is None or self.sheet_combo.currentIndex() < 0:
            return
        sheet_name = self.sheet_combo.currentText()
        if self.path.suffix.lower() == ".xlsx":
            workbook = openpyxl.load_workbook(str(self.path), read_only=True, data_only=True)
            try:
                sheet = workbook[sheet_name]
                rows = min(self.MAX_ROWS, sheet.max_row or 0)
                columns = min(self.MAX_COLUMNS, sheet.max_column or 0)
                values: list[list[object]] = []
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=rows,
                    min_col=1,
                    max_col=columns,
                    values_only=True,
                ):
                    values.append(list(row))
                total_rows, total_columns = sheet.max_row or 0, sheet.max_column or 0
            finally:
                workbook.close()
        else:
            workbook = xlrd.open_workbook(str(self.path), on_demand=True)
            try:
                sheet = workbook.sheet_by_name(sheet_name)
                rows = min(self.MAX_ROWS, sheet.nrows)
                columns = min(self.MAX_COLUMNS, sheet.ncols)
                values = [sheet.row_values(row, 0, columns) for row in range(rows)]
                total_rows, total_columns = sheet.nrows, sheet.ncols
            finally:
                workbook.release_resources()
        self.table.setRowCount(rows)
        self.table.setColumnCount(columns)
        self.table.setHorizontalHeaderLabels([get_column_letter(i) for i in range(1, columns + 1)])
        for row_index, row in enumerate(values):
            for column_index, value in enumerate(row):
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                self.table.setItem(row_index, column_index, QTableWidgetItem("" if value is None else str(value)))
        self.table.resizeColumnsToContents()
        self.info_label.setText(
            f"{total_rows} Zeilen · {total_columns} Spalten · angezeigt bis {rows} × {columns}"
        )
        if total_rows > self.MAX_ROWS or total_columns > self.MAX_COLUMNS:
            self.info_label.setText(
                self.info_label.text()
                + " · Hinweis: Anzeige ist auf 500 Zeilen und 50 Spalten begrenzt."
            )

    def find_next(self):
        query = self.search_input.text().strip().casefold()
        if not query:
            return
        start = self.table.currentRow() * max(1, self.table.columnCount()) + self.table.currentColumn() + 1
        cells = self.table.rowCount() * self.table.columnCount()
        for step in range(cells):
            index = (start + step) % cells
            row, column = divmod(index, self.table.columnCount())
            item = self.table.item(row, column)
            if item and query in item.text().casefold():
                self.table.setCurrentItem(item)
                self.table.scrollToItem(item)
                return
