from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import time
import unittest

import openpyxl
from docx import Document
from PyPDF2 import PdfWriter
from PySide6.QtCore import QPoint, QPointF, Qt, QtMsgType, qInstallMessageHandler
from PySide6.QtGui import QPixmap, QResizeEvent, QWheelEvent
from PySide6.QtWidgets import QApplication
from unittest.mock import Mock, patch

from app.gui.viewer import FileViewer, ImageViewerWidget
from app.gui.viewers import PdfViewerWidget, SpreadsheetViewerWidget, TextViewerWidget
from app.gui.workers.file_conversion_worker import FileConversionWorker


class ViewerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def process_until(self, predicate, timeout: float = 5):
        deadline = time.monotonic() + timeout
        while not predicate() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(0.01)
        self.app.processEvents()

    def test_pdf_text_and_multisheet_excel_controls(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "test.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=400)
            with pdf_path.open("wb") as handle:
                writer.write(handle)

            excel_path = root / "test.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "Übersicht"
            workbook.active["A1"] = "Musterwert"
            workbook.create_sheet("Details")["B2"] = "Treffer"
            workbook.save(excel_path)

            text_path = root / "test.txt"
            text_path.write_text("eins zwei drei", encoding="utf-8")

            viewer = FileViewer()
            viewer.open_file(pdf_path)
            self.assertEqual(viewer.pdf_viewer.document.pageCount(), 1)
            viewer.open_file(excel_path)
            self.assertEqual(viewer.spreadsheet_viewer.sheet_combo.count(), 2)
            viewer.spreadsheet_viewer.sheet_combo.setCurrentText("Details")
            self.assertEqual(viewer.spreadsheet_viewer.table.item(1, 1).text(), "Treffer")
            viewer.open_file(text_path)
            viewer.text_viewer.search_input.setText("zwei")
            viewer.text_viewer.find_next()
            self.assertTrue(viewer.text_viewer.editor.textCursor().hasSelection())
            viewer.shutdown()
            viewer.close()

    def test_spreadsheet_viewer_displays_explicit_truncation_hint(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            excel_path = root / "big.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Daten"
            for row in range(1, 505):
                for column in range(1, 52):
                    sheet.cell(row=row, column=column, value=f"{row}-{column}")
            workbook.save(excel_path)

            viewer = FileViewer()
            viewer.open_file(excel_path)

            self.assertIn("angezeigt bis 500 × 50", viewer.spreadsheet_viewer.info_label.text())
            self.assertIn("Hinweis: Anzeige ist auf 500 Zeilen und 50 Spalten begrenzt", viewer.spreadsheet_viewer.info_label.text())
            viewer.shutdown()
            viewer.close()

    def test_pdf_close_document_does_not_emit_nullptr_warning(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "warncheck.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=400)
            with pdf_path.open("wb") as handle:
                writer.write(handle)

            messages: list[str] = []

            def collector(message_type, _context, message):
                if message_type in {QtMsgType.QtWarningMsg, QtMsgType.QtCriticalMsg}:
                    messages.append(str(message))

            previous_handler = qInstallMessageHandler(collector)
            try:
                viewer = FileViewer()
                viewer.open_file(pdf_path)
                baseline = sum(
                    "invalid nullptr parameter" in message for message in messages
                )
                viewer.pdf_viewer.close_document()
                self.app.processEvents()
                viewer.shutdown()
                viewer.close()
            finally:
                qInstallMessageHandler(previous_handler)

            after_close = sum(
                "invalid nullptr parameter" in message for message in messages
            )
            self.assertEqual(after_close, baseline)

    def test_docx_preview_converts_to_pdf_asynchronously(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "preview.docx"
            document = Document()
            document.add_paragraph("Formatierter Inhalt")
            document.save(docx_path)
            pdf_path = root / "preview.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=400)
            with pdf_path.open("wb") as handle:
                writer.write(handle)

            def convert_to_pdf(*_args, **_kwargs):
                time.sleep(0.05)
                return pdf_path

            with patch(
                "app.services.document_converter.DocumentConverter.convert_word_to_pdf",
                side_effect=convert_to_pdf,
            ):
                viewer = FileViewer()
                viewer.open_file(docx_path)

                self.assertTrue(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.loading_widget)

                self.process_until(
                    lambda: viewer.stack.currentWidget() is viewer.pdf_viewer
                )
                self.process_until(lambda: not viewer.is_converting, timeout=1)

                self.assertFalse(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.pdf_viewer)
                self.assertEqual(viewer.pdf_viewer.document.pageCount(), 1)
                viewer.shutdown()
                viewer.close()

    def test_docx_preview_falls_back_to_text_when_pdf_conversion_fails(self):
        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "fallback.docx"
            document = Document()
            document.add_paragraph("Fallback Inhalt")
            document.save(docx_path)

            with patch(
                "app.services.document_converter.DocumentConverter.convert_word_to_pdf",
                side_effect=RuntimeError("kein Konverter"),
            ):
                viewer = FileViewer()
                viewer.open_file(docx_path)

                self.process_until(
                    lambda: viewer.stack.currentWidget() is viewer.text_viewer
                )
                self.process_until(lambda: not viewer.is_converting, timeout=1)

                self.assertFalse(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.text_viewer)
                text = viewer.text_viewer.editor.toPlainText()
                self.assertIn("Formatierte Word-Vorschau konnte nicht erstellt werden", text)
                self.assertIn("Fallback Inhalt", text)
                viewer.shutdown()
                viewer.close()

    def test_legacy_doc_processing_runs_asynchronously_with_loading_view(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.doc"
            path.write_bytes(
                b"Legacy customer document with enough readable text "
                b"for the Python fallback extraction."
            )

            with patch(
                "app.services.document_converter.DocumentConverter.convert_word_to_pdf",
                side_effect=RuntimeError("kein Konverter"),
            ):
                viewer = FileViewer()
                viewer.open_file(path)

                self.assertTrue(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.loading_widget)
                self.assertTrue(viewer.loading_indicator.is_running())

                self.process_until(
                    lambda: viewer.stack.currentWidget() is viewer.text_viewer
                )
                self.process_until(lambda: not viewer.is_converting, timeout=1)

                self.assertFalse(viewer.is_converting)
                self.assertIs(viewer.stack.currentWidget(), viewer.text_viewer)
                self.assertIn(
                    "Legacy customer document",
                    viewer.text_viewer.editor.toPlainText(),
                )
                viewer.shutdown()
                viewer.close()

    def test_image_viewer_scaling_and_control_wheel_edges(self):
        viewer = ImageViewerWidget()
        viewer.resizeEvent(QResizeEvent(viewer.size(), viewer.size()))
        viewer._update_display_pixmap()
        viewer.set_image(QPixmap())
        pixmap = QPixmap(20, 10)
        pixmap.fill(Qt.red)
        viewer.set_image(pixmap)
        self.assertFalse(viewer._image_label.pixmap().isNull())
        for delta in (120, -120, 0):
            event = Mock()
            event.modifiers.return_value = Qt.ControlModifier
            event.angleDelta.return_value = SimpleNamespace(y=lambda: delta)
            viewer.wheelEvent(event)
            event.accept.assert_called_once_with()
        viewer.resizeEvent(QResizeEvent(viewer.size(), viewer.size()))
        wheel = QWheelEvent(
            QPointF(1, 1), QPointF(1, 1), QPoint(), QPoint(0, 120),
            Qt.NoButton, Qt.NoModifier, Qt.NoScrollPhase, False,
        )
        viewer.wheelEvent(wheel)
        viewer.close()

    def test_file_routing_errors_images_unknown_and_external_actions(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            viewer = FileViewer()
            with patch("app.gui.viewer.QMessageBox.warning") as warning:
                viewer.open_file(root / "missing.txt")
            warning.assert_called_once()

            unknown = root / "file.bin"
            unknown.write_bytes(b"data")
            viewer.open_file(unknown)
            self.assertIn("keine interne Vorschau", viewer.text_viewer.editor.toPlainText())
            image = root / "image.png"
            pixmap = QPixmap(4, 4)
            pixmap.fill(Qt.blue)
            pixmap.save(str(image))
            viewer.open_file(image)
            self.assertIs(viewer.stack.currentWidget(), viewer.image_viewer)
            broken = root / "broken.png"
            broken.write_bytes(b"bad")
            viewer.open_file(broken)
            self.assertIn("Bild konnte nicht", viewer.text_viewer.editor.toPlainText())

            viewer.current_file = None
            viewer.open_externally()
            viewer.open_containing_folder()
            viewer.current_file = unknown
            with (
                patch("app.gui.viewer.QDesktopServices.openUrl", return_value=False),
                patch("app.gui.viewer.QMessageBox.warning") as warning,
            ):
                viewer.open_externally()
                viewer.open_containing_folder()
            self.assertEqual(warning.call_count, 2)
            with patch("app.gui.viewer.QDesktopServices.openUrl", return_value=True):
                viewer.open_externally()
                viewer.open_containing_folder()
            viewer.shutdown()

    def test_spreadsheet_and_conversion_callback_error_paths(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            viewer = FileViewer()
            xls = root / "legacy.xls"
            xls.write_bytes(b"data")
            with (
                patch.object(viewer.spreadsheet_viewer, "load", side_effect=ValueError),
                patch.object(viewer, "_start_conversion") as start,
            ):
                viewer._show_spreadsheet(xls)
            start.assert_called_once()
            xlsx = root / "modern.xlsx"
            with (
                patch.object(viewer.spreadsheet_viewer, "load", side_effect=ValueError),
                self.assertRaises(ValueError),
            ):
                viewer._show_spreadsheet(xlsx)

            viewer._load_generation = 2
            viewer._on_conversion_completed(1, "ignored", None, {}, "")
            with patch.object(viewer, "_handle_conversion_error") as error:
                viewer._on_conversion_completed(2, "operation", None, {}, "bad")
            error.assert_called_once_with("operation", "bad")
            with patch.object(viewer, "_handle_conversion_error") as error:
                viewer._on_conversion_completed(2, "operation", None, {}, "abgebrochen")
            error.assert_not_called()

            converted = root / "converted.xlsx"
            converted.touch()
            with patch.object(viewer.spreadsheet_viewer, "load"):
                viewer._on_conversion_completed(
                    2, FileConversionWorker.XLS_TO_XLSX, converted,
                    {"converted": True, "tool": "LibreOffice"}, "",
                )
            self.assertIn("LibreOffice", viewer.conversion_meta_label.text())
            pdf = root / "converted.pdf"
            pdf.touch()
            with patch.object(viewer.pdf_viewer, "load"):
                viewer._on_conversion_completed(
                    2, FileConversionWorker.WORD_TO_PDF, pdf, {}, "",
                )
            viewer._word_preview_error = "format"
            viewer._on_conversion_completed(
                2, FileConversionWorker.EXTRACT_DOC, "legacy text", {}, "",
            )
            self.assertIn("format", viewer.text_viewer.editor.toPlainText())
            viewer._on_conversion_completed(
                2, FileConversionWorker.EXTRACT_DOC, "plain text", {}, "",
            )
            viewer._on_conversion_completed(2, "unknown", None, {}, "")
            viewer._on_conversion_completed(
                2, FileConversionWorker.EXTRACT_DOC, "", {}, "",
            )
            self.assertIn("konnte nicht", viewer.text_viewer.editor.toPlainText())

            def change_generation_and_fail(_path):
                viewer._load_generation += 1
                raise ValueError("changed")

            with patch.object(
                viewer.spreadsheet_viewer, "load",
                side_effect=change_generation_and_fail,
            ):
                viewer._on_conversion_completed(
                    viewer._load_generation,
                    FileConversionWorker.XLS_TO_XLSX, converted, {}, "",
                )
            viewer.shutdown()

    def test_conversion_error_fallbacks_worker_cleanup_and_shutdown(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            viewer = FileViewer()
            docx = root / "bad.docx"
            docx.touch()
            viewer.current_file = docx
            with (
                patch.object(viewer, "_show_docx_text", side_effect=ValueError("bad fallback")),
            ):
                viewer._handle_conversion_error(FileConversionWorker.WORD_TO_PDF, "convert")
            self.assertIn("bad fallback", viewer.text_viewer.editor.toPlainText())
            doc = root / "old.doc"
            viewer.current_file = doc
            with patch.object(viewer, "_show_legacy_doc") as legacy:
                viewer._handle_conversion_error(FileConversionWorker.WORD_TO_PDF, "convert")
            legacy.assert_called_once_with(doc)
            viewer.current_file = None
            viewer._handle_conversion_error("other", "failed")
            viewer.current_file = root / "other.txt"
            viewer._handle_conversion_error(FileConversionWorker.WORD_TO_PDF, "failed")

            worker = Mock()
            worker.isRunning.return_value = True
            viewer._conversion_workers.add(worker)
            self.assertTrue(viewer.is_converting)
            viewer._cancel_conversions()
            worker.requestInterruption.assert_called_once_with()
            viewer._release_conversion_worker(worker)
            worker.deleteLater.assert_called_once_with()
            converter = Mock()
            viewer._active_converter = converter
            viewer._cleanup_active_converter()
            converter.cleanup.assert_called_once_with()

            stopped = Mock()
            stopped.isRunning.return_value = False
            viewer._conversion_workers.add(stopped)
            viewer.shutdown()
            stopped.wait.assert_called_once_with()
            stopped.cleanup.assert_called_once_with()

    def test_focused_viewer_control_and_legacy_spreadsheet_edges(self):
        pdf = PdfViewerWidget()
        with TemporaryDirectory() as directory:
            invalid = Path(directory) / "invalid.pdf"
            invalid.write_text("bad", encoding="utf-8")
            with self.assertRaises(ValueError):
                pdf.load(invalid)
        document = Mock()
        document.pageCount.return_value = 3
        navigator = Mock()
        navigator.currentPage.return_value = 1
        navigator.currentZoom.return_value = 1.0
        view = Mock()
        view.pageNavigator.return_value = navigator
        view.zoomFactor.return_value = 10
        pdf.document = document
        pdf.view = view
        pdf._change_page(1)
        navigator.jump.assert_called_once()
        pdf._zoom(2)
        view.setZoomFactor.assert_called_with(5.0)
        pdf.search_model = Mock()
        pdf.search_input.setText("")
        pdf.find_next()
        pdf.search_input.setText("query")
        pdf.search_model.count.return_value = 0
        pdf.find_next()
        pdf.search_model.count.return_value = 2
        view.currentSearchResultIndex.return_value = 1
        pdf.find_next()
        view.setCurrentSearchResultIndex.assert_called_with(0)
        real_pdf = PdfViewerWidget()
        with patch("app.gui.viewers.pdf_viewer.QApplication.instance", return_value=None):
            real_pdf.close_document()

        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "table.docx"
            document_file = Document()
            document_file.add_paragraph("paragraph")
            table = document_file.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "left"
            table.cell(0, 1).text = "right"
            document_file.save(docx_path)
            file_viewer = FileViewer()
            file_viewer._show_docx_text(docx_path, prefix="prefix")
            self.assertIn("left\tright", file_viewer.text_viewer.editor.toPlainText())
            file_viewer._show_docx_text(docx_path)

        conversion_worker = FileConversionWorker(
            2, FileConversionWorker.XLS_TO_XLSX, Path("source.xls"), parent=file_viewer
        )
        conversion_worker.converter = Mock()
        with (
            patch.object(FileViewer, "sender", return_value=conversion_worker),
            patch.object(file_viewer.spreadsheet_viewer, "load"),
        ):
            file_viewer._load_generation = 2
            file_viewer._on_conversion_completed(
                2, FileConversionWorker.XLS_TO_XLSX, Path("result.xlsx"), {}, ""
            )
        file_viewer.shutdown()

        spreadsheet = SpreadsheetViewerWidget()
        spreadsheet._load_current_sheet()
        workbook = Mock()
        workbook.sheet_names.return_value = ["Legacy"]
        sheet = Mock(nrows=1, ncols=2)
        sheet.row_values.return_value = [1.0, None]
        workbook.sheet_by_name.return_value = sheet
        with patch("app.gui.viewers.spreadsheet_viewer.xlrd.open_workbook", return_value=workbook):
            spreadsheet.load(Path("legacy.xls"))
        workbook.release_resources.assert_called()
        self.assertEqual(spreadsheet.table.item(0, 0).text(), "1")
        spreadsheet.search_input.setText("")
        spreadsheet.find_next()
        spreadsheet.search_input.setText("absent")
        spreadsheet.find_next()
        spreadsheet.search_input.setText("1")
        spreadsheet.find_next()
        self.assertEqual(spreadsheet.table.currentItem().text(), "1")

        text = TextViewerWidget()
        text.set_text("one two")
        text.search_input.setText("")
        text.find_next()
        text.search_input.setText("two")
        text.find_next()
        text.find_next()
        self.assertTrue(text.editor.textCursor().hasSelection())


if __name__ == "__main__":
    unittest.main()
